from datetime import datetime, date
import os
import time

import requests
import numpy as np

from core.qt_runtime import configure_qt_runtime
configure_qt_runtime()

from PySide6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QLabel, QTableWidget, QTableWidgetItem,
    QHeaderView, QAbstractItemView, QComboBox, QPushButton, QMessageBox, QTabWidget,
    QLineEdit, QDialog, QScrollArea, QFrame, QDialogButtonBox, QFileDialog, QMenu,
    QDoubleSpinBox, QSpinBox, QDateEdit, QFormLayout, QStackedWidget, QGridLayout,
    QInputDialog, QCheckBox,
)
from PySide6.QtCore import Qt, Signal, QTimer, QDate, QPropertyAnimation, QThread
from PySide6.QtGui import QColor, QBrush

import matplotlib
matplotlib.use('QtAgg')
from matplotlib.backends.backend_qtagg import FigureCanvasQTAgg as FigureCanvas
from matplotlib.figure import Figure

from core.constants import DARK_BG, CHART_BG, MUTED, RED, GREEN
from core.data_worker import DataWorker, is_trading_time
from core.cache import (
    save_panel_cache, load_panel_cache, save_watchlist, load_watchlist,
    save_watchlist_fund_history, load_watchlist_fund_history,
    save_watchlist_groups, load_watchlist_groups,
    classify_code, compute_buy, compute_sell,
    load_option_watchlist, save_option_watchlist,
    compute_option_metrics, option_days_to_expiry,
    add_stock_to_watchlist,
    get_default_position_value,
)
from core.net_setup import push2_get
from core.data_source import quotes_routed, is_etf_code
from core.option_quote_provider import fetch_option_quotes
from core.portfolio_data import add_holding
from ui.widgets.table_utils import configure_no_truncation


OPTION_AUTO_REFRESH_INTERVAL_MS = 8000


def is_option_auto_refresh_time(value=None) -> bool:
    current = value if value is not None else datetime.now().time()
    if isinstance(current, datetime):
        current = current.time()
    minute = current.hour * 60 + current.minute
    return (
        (9 * 60 + 25) <= minute <= (11 * 60 + 35)
        or (12 * 60 + 55) <= minute <= (15 * 60 + 5)
    )


_PERIODS = {
    '今日': {
        'fid': 'f62',
        'fields': 'f12,f14,f2,f3,f8,f10,f6,f62,f184,f66,f69,f72,f75,f78,f81,f84,f87,f204,f205,f124,f100',
        'pct': 'f3',
        'main': ('f62', 'f184'),
        'super': ('f66', 'f69'),
        'big': ('f72', 'f75'),
        'mid': ('f78', 'f81'),
        'small': ('f84', 'f87'),
    },
    '3日': {
        'fid': 'f267',
        'fields': 'f12,f14,f2,f127,f8,f10,f6,f267,f268,f269,f270,f271,f272,f273,f274,f275,f276,f257,f258,f124,f100',
        'pct': 'f127',
        'main': ('f267', 'f268'),
        'super': ('f269', 'f270'),
        'big': ('f271', 'f272'),
        'mid': ('f273', 'f274'),
        'small': ('f275', 'f276'),
    },
    '5日': {
        'fid': 'f164',
        'fields': 'f12,f14,f2,f109,f8,f10,f6,f164,f165,f166,f167,f168,f169,f170,f171,f172,f173,f257,f258,f124,f100',
        'pct': 'f109',
        'main': ('f164', 'f165'),
        'super': ('f166', 'f167'),
        'big': ('f168', 'f169'),
        'mid': ('f170', 'f171'),
        'small': ('f172', 'f173'),
    },
    '10日': {
        'fid': 'f174',
        'fields': 'f12,f14,f2,f160,f8,f10,f6,f174,f175,f176,f177,f178,f179,f180,f181,f182,f183,f260,f261,f124,f100',
        'pct': 'f160',
        'main': ('f174', 'f175'),
        'super': ('f176', 'f177'),
        'big': ('f178', 'f179'),
        'mid': ('f180', 'f181'),
        'small': ('f182', 'f183'),
    },
}

_HEADERS = ['排名', '代码', '名称', '板块', '最新价', '涨跌幅', '换手率', '量比', '成交额', '主力净额', '主力占比', '超大单', '大单', '中单', '小单', '操作']

_HEADER_TIPS = {
    '涨跌幅': '当前周期涨跌幅：今日/3日/5日/10日随右上角周期切换变化',
    '主力净额': '主力净流入金额 = 超大单净额 + 大单净额；红色为净流入，绿色为净流出',
    '主力占比': '主力净额 / 成交额，反映主力资金相对成交量的强弱',
    '换手率': '当日成交股数 / 流通股本；越高代表筹码交换越活跃',
    '量比': '当前成交量相对近5日同时间平均成交量的倍数；>1.2 放量，<0.8 缩量',
    '成交额': '当日累计成交金额；换手率、量比、主力占比都需要结合成交额一起看',
    '超大单': '超大单净额，通常代表更大资金单子的净买卖方向',
    '大单': '大单净额，和超大单一起构成主力净额',
    '中单': '中单净额，不计入主力净额',
    '小单': '小单净额，常被用来观察散户侧资金方向，但不能简单等同于散户真实行为',
}

_LINE_EDIT_STYLE = (
    'QLineEdit { background-color: #1d2d4a; color: #ddd; border: 1px solid #4a6a8e; '
    'border-radius: 4px; padding: 4px 8px; font-size: 12px; }'
)
_COMBO_STYLE = (
    'QComboBox { background-color: #1d2d4a; color: #ddd; border: 1px solid #4a6a8e; '
    'border-radius: 4px; padding: 4px 8px; font-size: 12px; }'
    'QComboBox QAbstractItemView { background-color: #1d2d4a; color: #ddd; '
    'selection-background-color: #2a3a5e; }'
)
_SPIN_STYLE = (
    'QDoubleSpinBox, QSpinBox, QDateEdit { background-color: #1d2d4a; color: #ddd; '
    'border: 1px solid #4a6a8e; border-radius: 4px; padding: 4px 8px; font-size: 12px; }'
)
_NUMERIC_SPIN_STYLE = (
    _SPIN_STYLE +
    '''
    QSpinBox, QDoubleSpinBox { padding-right: 34px; }
    QSpinBox::up-button, QDoubleSpinBox::up-button {
        subcontrol-origin: border;
        subcontrol-position: top right;
        width: 28px;
        border-left: 1px solid #4a6a8e;
        border-bottom: 1px solid #2a4260;
        border-top-right-radius: 4px;
        background-color: #223957;
    }
    QSpinBox::down-button, QDoubleSpinBox::down-button {
        subcontrol-origin: border;
        subcontrol-position: bottom right;
        width: 28px;
        border-left: 1px solid #4a6a8e;
        border-bottom-right-radius: 4px;
        background-color: #182b46;
    }
    QSpinBox::up-button:hover, QDoubleSpinBox::up-button:hover,
    QSpinBox::down-button:hover, QDoubleSpinBox::down-button:hover {
        background-color: #2d4c72;
    }
    '''
)


def _apply_numeric_spin_style(spin):
    spin.setStyleSheet(_NUMERIC_SPIN_STYLE)
    spin.setAccelerated(True)


def _to_float(value, default=0.0):
    try:
        if value in ('-', None):
            return default
        return float(value)
    except (TypeError, ValueError):
        return default


def _format_money(value):
    value = _to_float(value)
    if abs(value) >= 100000000:
        return f'{value / 100000000:.2f}亿'
    if abs(value) >= 10000:
        return f'{value / 10000:.0f}万'
    return f'{value:.0f}'


def _format_pct(value, signed=False):
    value = _to_float(value)
    sign = '+' if signed and value >= 0 else ''
    return f'{sign}{value:.2f}%'


def _format_price(value, code=None, is_option=False):
    """价格字符串：ETF/期权 4 位小数，其余 3 位；空值返回 '—'。"""
    value = _to_float(value)
    if not value:
        return '—'
    from core.data_source import fmt_price
    return fmt_price(value, code, is_option=is_option)


def _default_wl_shares(price) -> int | None:
    """自选默认股数 = ROUND(默认市值/现价, -2)，凑整到 100 股。"""
    price = _to_float(price)
    if price <= 0:
        return None
    mv = float(get_default_position_value())
    return int((mv / price + 50) // 100 * 100)


def _format_volume_ratio(value):
    value = _to_float(value)
    if value >= 1.2:
        label = '放量'
    elif value <= 0.8:
        label = '缩量'
    else:
        label = '平量'
    return f'{value:.2f} {label}'


def _format_time(value):
    try:
        return datetime.fromtimestamp(int(value)).strftime('%H:%M:%S')
    except (TypeError, ValueError, OSError):
        return datetime.now().strftime('%H:%M:%S')


def _snapshot_date():
    return datetime.now().strftime('%Y-%m-%d')


def _history_cache_name(direction, period):
    return f"stock_history_{direction}_{period}"


def _save_stock_history_snapshot(direction, snap, max_days=10):
    if not snap or not snap.get('rows'):
        return
    cache_name = _history_cache_name(direction, snap.get('period') or '今日')
    dataset = load_panel_cache(cache_name) or {}
    records = dataset.get('records') or []
    date_key = snap.get('date') or _snapshot_date()
    snap = dict(snap)
    snap['date'] = date_key
    records = [r for r in records if r.get('date') != date_key]
    records.append(snap)
    records.sort(key=lambda r: r.get('date', ''), reverse=True)
    save_panel_cache(cache_name, {'records': records[:max_days]})


def _load_latest_stock_history_snapshot(direction, period):
    dataset = load_panel_cache(_history_cache_name(direction, period)) or {}
    records = [r for r in (dataset.get('records') or []) if r.get('rows')]
    if records:
        records.sort(key=lambda r: (r.get('date', ''), r.get('fetch_time', '')), reverse=True)
        return records[0]
    return load_panel_cache(f"stock_{direction}_{period}")


def _enrich_rows_with_tencent(rows):
    """从腾讯刷新 rows 的价格/涨跌幅/换手率/成交额。

    返回新的 rows 列表（不修改原列表）。
    仅覆盖腾讯能提供的字段；main_net / super_net 等东财独占字段保持缓存原值。
    失败时抛错，由调用者决定是否放弃。
    """
    codes = [r.get('code') for r in rows if r.get('code')]
    if not codes:
        return list(rows)
    quotes = quotes_routed(codes, timeout=6)
    out = []
    for r in rows:
        nr = dict(r)
        q = quotes.get(r.get('code'), {})
        if q:
            if q.get('name'):
                nr['name'] = q['name']
            if q.get('price'):
                nr['price'] = q['price']
            if q.get('pct') is not None:
                nr['pct'] = q['pct']
            if q.get('turnover'):
                nr['turnover'] = q['turnover']
            if q.get('amount'):
                nr['amount'] = q['amount']
        out.append(nr)
    return out


def _fetch_stock_top(period='今日', top_n=20, ascending=False):
    conf = _PERIODS[period]
    params = {
        'fid': conf['fid'],
        'po': '0' if ascending else '1',
        'pz': str(top_n),
        'pn': '1',
        'np': '1',
        'fltt': '2',
        'invt': '2',
        'ut': 'b2884a393a59ad64002292a3e90d46a5',
        'fs': 'm:0+t:6+f:!2,m:0+t:13+f:!2,m:0+t:80+f:!2,m:1+t:2+f:!2,m:1+t:23+f:!2,m:0+t:7+f:!2,m:1+t:3+f:!2',
        'fields': conf['fields'],
    }
    try:
        r = push2_get('/api/qt/clist/get', params=params, timeout=8)
        r.raise_for_status()
        data = r.json().get('data') or {}
        diff = data.get('diff') or []
    except Exception as primary_error:
        # 兜底：Tencent 刷新最近一次缓存的 rows（价格/涨跌幅/换手/成交额）
        # 主力净额等东财独占字段保持缓存值（会偏陈但总比空表好）
        direction = 'out' if ascending else 'in'
        cached = _load_latest_stock_history_snapshot(direction, period)
        if not cached or not cached.get('rows'):
            raise primary_error
        try:
            enriched = _enrich_rows_with_tencent(cached['rows'])
        except Exception:
            raise primary_error
        return {
            'date': cached.get('date') or _snapshot_date(),
            'period': period,
            'rows': enriched,
            'total': cached.get('total') or len(enriched),
            'quote_time': '腾讯实时',
            'fetch_time': datetime.now().strftime('%H:%M:%S'),
            'fallback': True,
        }

    rows = []
    latest_ts = None
    for i, item in enumerate(diff[:top_n], 1):
        latest_ts = item.get('f124') or latest_ts
        main_net, main_pct = conf['main']
        super_net, super_pct = conf['super']
        big_net, big_pct = conf['big']
        mid_net, mid_pct = conf['mid']
        small_net, small_pct = conf['small']
        rows.append({
            'rank': i,
            'code': str(item.get('f12') or ''),
            'name': str(item.get('f14') or ''),
            'sector': str(item.get('f100') or ''),
            'price': item.get('f2'),
            'pct': item.get(conf['pct']),
            'turnover': item.get('f8'),
            'volume_ratio': item.get('f10'),
            'amount': item.get('f6'),
            'main_net': item.get(main_net),
            'main_pct': item.get(main_pct),
            'super_net': item.get(super_net),
            'super_pct': item.get(super_pct),
            'big_net': item.get(big_net),
            'big_pct': item.get(big_pct),
            'mid_net': item.get(mid_net),
            'mid_pct': item.get(mid_pct),
            'small_net': item.get(small_net),
            'small_pct': item.get(small_pct),
        })

    return {
        'date': _snapshot_date(),
        'period': period,
        'rows': rows,
        'total': int(data.get('total') or len(rows)),
        'quote_time': _format_time(latest_ts),
        'fetch_time': datetime.now().strftime('%H:%M:%S'),
    }


_SINA_GAIN_URL = (
    'https://vip.stock.finance.sina.com.cn/quotes_service/api/json_v2.php/'
    'Market_Center.getHQNodeDataSimple'
)
_SINA_GAIN_HEADERS = {
    'User-Agent': (
        'Mozilla/5.0 (Windows NT 10.0; Win64; x64) '
        'AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36'
    ),
    'Referer': 'https://finance.sina.com.cn/',
}


def _fetch_stock_gain_sina(top_n=30, ascending=False):
    """新浪全市场涨幅/跌幅新鲜榜单（push2 熔断时的二级兜底）。

    覆盖沪深 A 股；量比字段新浪不提供，填 0。amount 单位：元。
    """
    params = {
        'page': '1',
        'num': str(top_n),
        'sort': 'changepercent',
        'asc': '1' if ascending else '0',
        'node': 'hs_a',
        '_s_r_a': 'init',
    }
    resp = requests.get(
        _SINA_GAIN_URL,
        params=params,
        headers=_SINA_GAIN_HEADERS,
        timeout=8,
    )
    resp.raise_for_status()
    items = resp.json()
    if not isinstance(items, list):
        raise ValueError(f'新浪涨跌幅榜格式异常: {type(items)}')
    rows = []
    for i, item in enumerate(items[:top_n], 1):
        sym = str(item.get('symbol', ''))
        code = sym[2:] if sym.startswith(('sh', 'sz', 'bj')) else sym
        try:
            pct = float(item.get('changepercent') or 0)
        except (TypeError, ValueError):
            pct = 0.0
        try:
            price = float(item.get('trade') or 0)
        except (TypeError, ValueError):
            price = 0.0
        try:
            turnover = float(item.get('turnoverratio') or 0)
        except (TypeError, ValueError):
            turnover = 0.0
        try:
            amount = float(item.get('amount') or 0)
        except (TypeError, ValueError):
            amount = 0.0
        rows.append({
            'rank': i,
            'code': code,
            'name': str(item.get('name', '')),
            'sector': '',
            'price': price,
            'pct': pct,
            'turnover': turnover,
            'volume_ratio': 0.0,
            'amount': amount,
        })
    if not rows:
        raise ValueError('新浪涨跌幅榜返回空数据')
    return rows


def _fetch_stock_gain_top(top_n=20, ascending=False):
    params = {
        'fid': 'f3',
        'po': '0' if ascending else '1',
        'pz': str(top_n),
        'pn': '1',
        'np': '1',
        'fltt': '2',
        'invt': '2',
        'ut': 'b2884a393a59ad64002292a3e90d46a5',
        'fs': 'm:0+t:6+f:!2,m:0+t:13+f:!2,m:0+t:80+f:!2,m:1+t:2+f:!2,m:1+t:23+f:!2,m:0+t:7+f:!2,m:1+t:3+f:!2',
        'fields': 'f12,f14,f2,f3,f8,f10,f6,f62,f184,f66,f69,f72,f75,f78,f81,f84,f87,f124,f100',
    }
    direction = 'drop' if ascending else 'gain'
    try:
        r = push2_get('/api/qt/clist/get', params=params, timeout=8)
        r.raise_for_status()
        data = r.json().get('data') or {}
        diff = data.get('diff') or []
    except Exception as primary_error:
        # 兆底①：新浪新鲜涨跌幅榜（全市场，不受旧缓存范围限制）
        try:
            sina_rows = _fetch_stock_gain_sina(top_n=top_n, ascending=ascending)
            return {
                'date': _snapshot_date(),
                'direction': direction,
                'rows': sina_rows,
                'total': len(sina_rows),
                'quote_time': '新浪实时',
                'fetch_time': datetime.now().strftime('%H:%M:%S'),
                'fallback': True,
            }
        except Exception:
            pass
        # 兆底②：Tencent 刷新缓存 rows 的价格/涨跌幅，以 fresh pct 重新排序
        cached = load_panel_cache(f'stock_{direction}')
        if not cached or not cached.get('rows'):
            raise primary_error
        try:
            enriched = _enrich_rows_with_tencent(cached['rows'])
        except Exception:
            raise primary_error
        # 涨幅/跌幅榜按 fresh pct 重新排序（涨幅高 -> 低；跌幅低 -> 高）
        enriched.sort(key=lambda r: r.get('pct') if r.get('pct') is not None else 0, reverse=not ascending)
        for i, row in enumerate(enriched, 1):
            row['rank'] = i
        return {
            'date': _snapshot_date(),
            'direction': direction,
            'rows': enriched,
            'total': cached.get('total') or len(enriched),
            'quote_time': '腾讯实时',
            'fetch_time': datetime.now().strftime('%H:%M:%S'),
            'fallback': True,
        }
    rows = []
    latest_ts = None
    for i, item in enumerate(diff[:top_n], 1):
        latest_ts = item.get('f124') or latest_ts
        rows.append({
            'rank': i,
            'code': str(item.get('f12') or ''),
            'name': str(item.get('f14') or ''),
            'sector': str(item.get('f100') or ''),
            'price': item.get('f2'),
            'pct': item.get('f3'),
            'turnover': item.get('f8'),
            'volume_ratio': item.get('f10'),
            'amount': item.get('f6'),
            'main_net': item.get('f62'),
            'main_pct': item.get('f184'),
            'super_net': item.get('f66'),
            'super_pct': item.get('f69'),
            'big_net': item.get('f72'),
            'big_pct': item.get('f75'),
            'mid_net': item.get('f78'),
            'mid_pct': item.get('f81'),
            'small_net': item.get('f84'),
            'small_pct': item.get('f87'),
        })
    return {
        'date': _snapshot_date(),
        'direction': 'drop' if ascending else 'gain',
        'rows': rows,
        'total': int(data.get('total') or len(rows)),
        'quote_time': _format_time(latest_ts),
        'fetch_time': datetime.now().strftime('%H:%M:%S'),
    }


_WATCHLIST_DAILY_PNL_HEADER = '当日收益\n(现价-昨收)'
_WATCHLIST_DAILY_PNL_FORMULA = '当日收益 =（现价 − 昨收）× 数量'
_WATCHLIST_TOTAL_PNL_FORMULA = '当日总收益 = 当前分组个股当日收益合计 + 期权当日收益合计'
_WATCHLIST_HEADERS = ['代码', '名称', '数量', '加入价', '最新价', '昨收价', '涨跌幅', '主力净额(亿)', '主力净比(%)', '超大单净额', '大单净额', '累计盈亏', '盈亏比例', _WATCHLIST_DAILY_PNL_HEADER, '操作']
_WATCHLIST_HEADER_TIPS = {
    _WATCHLIST_DAILY_PNL_HEADER: _WATCHLIST_DAILY_PNL_FORMULA,
}

_TENCENT_QUOTE_URL = 'https://qt.gtimg.cn/q='


def _is_valid_stock_code(code: str) -> bool:
    """校验是否为可加入自选的 A 股/ETF 代码，避免把日期/金额等6位数字误判为代码。

    覆盖范围：深市主板(00xxxx)、创业板(30xxxx)、沪市主板(60xxxx)、
    科创板(68xxxx)、北交所(83xxxx-88xxxx)、B股(20xxxx/90xxxx)，
    以及自选页已支持报价/展示的 ETF/LOF 基金代码。
    """
    if len(code) != 6 or not code.isdigit():
        return False
    if is_etf_code(code):
        return True
    n = int(code)
    return (
        (0 <= n <= 4999)            # 深市主板 000001-004999
        or (300000 <= n <= 301999)  # 创业板
        or (600000 <= n <= 605999)  # 沪市主板
        or (688000 <= n <= 689999)  # 科创板
        or (830000 <= n <= 889999)  # 北交所
        or (920000 <= n <= 929999)  # 北交所新代码段
        or (200000 <= n <= 200999)  # 深市 B 股
        or (900900 <= n <= 900999)  # 沪市 B 股
    )


def _clean_code(code):
    """去掉 sh/sz 前缀，返回纯6位数字代码。"""
    code = code.strip()
    if code.startswith(('sh', 'sz')):
        return code[2:]
    return code


def _code_to_tencent(code):
    code = code.strip()
    if code.startswith(('sh', 'sz', 'bj')):
        return code
    if code.startswith('92'):
        return f'bj{code}'
    if code.startswith(('6', '5', '9')):
        return f'sh{code}'
    if code.startswith(('4', '8')):
        return f'bj{code}'
    return f'sz{code}'


def _code_to_secid(code):
    code = code.strip()
    if code.startswith('sh'):
        return f'1.{code[2:]}'
    if code.startswith('sz'):
        return f'0.{code[2:]}'
    if code.startswith('bj'):
        return f'0.{code[2:]}'
    if code.startswith('92'):
        return f'0.{code}'
    if code.startswith(('6', '5', '9')):
        return f'1.{code}'
    if code.startswith(('4', '8')):
        return f'0.{code}'
    return f'0.{code}'


def _fetch_watchlist_data(codes):
    if not codes:
        return {'rows': [], 'fetch_time': datetime.now().strftime('%H:%M:%S'), 'date': _snapshot_date()}

    # 1. 腾讯行情：名称/最新价/涨跌幅
    tc_codes = ','.join(_code_to_tencent(c) for c in codes)
    quotes = {}
    r = requests.get(
        _TENCENT_QUOTE_URL + tc_codes,
        headers={'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36'},
        timeout=10,
    )
    r.raise_for_status()
    text = r.content.decode('gbk', 'replace')
    for line in text.strip().split(';'):
        line = line.strip()
        if not line or '=' not in line:
            continue
        val = line.split('=', 1)[1].strip('"')
        fields = val.split('~')
        if len(fields) < 33:
            continue
        code = fields[2].strip()
        try:
            price = float(fields[3]) if fields[3] else 0.0
            pre_close = float(fields[4]) if fields[4] else 0.0
            pct = float(fields[32]) if fields[32] else 0.0
        except (ValueError, TypeError, IndexError):
            price, pre_close, pct = 0.0, 0.0, 0.0
        quotes[code] = {'name': fields[1], 'price': price, 'pct': pct, 'pre_close': pre_close}

    # 2. Push2 资金流：主力净额/主力净比/超大单/大单
    funds = {}
    try:
        secids = ','.join(_code_to_secid(c) for c in codes)
        params = {
            'fltt': '2',
            'invt': '2',
            'ut': 'b2884a393a59ad64002292a3e90d46a5',
            'fields': 'f12,f14,f62,f184,f66,f72',
            'secids': secids,
        }
        resp = push2_get('/api/qt/ulist.np/get', params=params, timeout=8)
        resp.raise_for_status()
        data = resp.json().get('data') or {}
        diff = data.get('diff') or []
        for item in diff:
            c = str(item.get('f12') or '')
            funds[c] = {
                'main_net': item.get('f62'),
                'main_pct': item.get('f184'),
                'super_net': item.get('f66'),
                'big_net': item.get('f72'),
            }
    except Exception:
        pass

    # 3. 合并（quotes/funds 的 key 是纯6位代码）
    rows = []
    for code in codes:
        clean = _clean_code(code)
        q = quotes.get(clean, {})
        f = funds.get(clean, {})
        rows.append({
            'code': code,
            'name': q.get('name', ''),
            'price': q.get('price', 0.0),
            'pre_close': q.get('pre_close', 0.0),
            'pct': q.get('pct', 0.0),
            'main_net': f.get('main_net'),
            'main_pct': f.get('main_pct'),
            'super_net': f.get('super_net'),
            'big_net': f.get('big_net'),
            'has_fund': bool(f),
        })

    return {
        'rows': rows,
        'fetch_time': datetime.now().strftime('%H:%M:%S'),
        'date': _snapshot_date(),
    }


class _WatchlistFundDialog(QDialog):
    """自选股近10日主力净额回放弹窗（红柱净流入，绿柱净流出）。"""

    def __init__(self, code, name, records, parent=None):
        super().__init__(parent)
        self.setWindowTitle(f'{name}({code}) 近10日主力净额')
        self.setMinimumSize(580, 380)
        self.resize(640, 420)
        self.setStyleSheet(f"background-color: {DARK_BG}; color: #ffffff;")

        layout = QVBoxLayout(self)
        layout.setContentsMargins(10, 10, 10, 10)

        fig = Figure(facecolor=DARK_BG, tight_layout=True)
        canvas = FigureCanvas(fig)
        layout.addWidget(canvas)

        ax = fig.add_subplot(111)
        ax.set_facecolor(CHART_BG)

        if not records:
            ax.text(0.5, 0.5, '暂无历史资金数据\n\n请先在交易时段刷新自选股',
                    transform=ax.transAxes, ha='center', va='center',
                    fontsize=13, color=MUTED, linespacing=1.8)
            ax.set_xticks([])
            ax.set_yticks([])
            canvas.draw_idle()
            return

        # 按日期正序排列
        sorted_recs = sorted(records, key=lambda r: r.get('date', ''))
        dates = [r['date'][5:] for r in sorted_recs]  # MM-DD
        main_nets = []
        for r in sorted_recs:
            v = r.get('main_net')
            try:
                main_nets.append(float(v) / 1e8 if v is not None else 0.0)
            except (TypeError, ValueError):
                main_nets.append(0.0)

        x = np.arange(len(dates))
        colors = [RED if v >= 0 else GREEN for v in main_nets]
        bars = ax.bar(x, main_nets, color=colors, width=0.55, edgecolor='none')

        for bar, val in zip(bars, main_nets):
            y = bar.get_height()
            va = 'bottom' if y >= 0 else 'top'
            ax.text(bar.get_x() + bar.get_width() / 2, y,
                    f'{val:.2f}', ha='center', va=va,
                    fontsize=8, color='#dddddd', fontweight='bold')

        ax.set_xticks(x)
        ax.set_xticklabels(dates, fontsize=9)
        ax.axhline(y=0, color='white', linewidth=0.5, alpha=0.3)
        ax.set_title(f'{name}({code}) 近{len(sorted_recs)}日主力净额（亿元）',
                     fontsize=13, fontweight='bold', color='white', pad=10)
        ax.set_ylabel('主力净额（亿元）', color=MUTED, fontsize=10)
        ax.tick_params(colors='white', labelsize=9)
        ax.grid(axis='y', color='#444444', linewidth=0.3, alpha=0.5)
        for spine in ax.spines.values():
            spine.set_color('#444444')

        canvas.draw_idle()


def _parse_stock_code(raw: str):
    """解析用户输入，返回纯6位数字代码，失败返回 None。
    支持：000001 / sh000001 / SH000001 / 000001.SZ / 000001.SH / 000001.CSI
    """
    s = raw.strip().upper()
    if not s:
        return None
    if '.' in s:
        s = s.split('.', 1)[0]
    for prefix in ('SH', 'SZ', 'BJ'):
        if s.startswith(prefix) and len(s) > 2:
            s = s[len(prefix):]
            break
    if s.isdigit() and len(s) == 6:
        return s
    return None


class AddWatchlistDialog(QDialog):
    """添加自选股弹窗：代码（必填）+ 数量（留空按默认市值估算）/成本价（可选）+ 分组。"""

    _NEW_GROUP_ITEM = '＋ 新建分组…'

    def __init__(self, existing_by_group: dict | None = None, parent=None,
                 prefill_code: str = '', prefill_name: str = '', prefill_price: float = 0.0,
                 groups: list[str] | None = None, default_group: str = ''):
        super().__init__(parent)
        self.setWindowTitle('添加自选股')
        self.setFixedSize(340, 540)
        self.setStyleSheet(f'background-color: {DARK_BG}; color: #ffffff;')
        self._existing_by_group = existing_by_group or {}
        self.result_code = None
        self.result_name = ''
        self.result_shares = None
        self.result_cost_price = None
        self.result_theme = ''
        self.result_buy_date = ''
        self.result_group = default_group or (groups[0] if groups else '')

        layout = QVBoxLayout(self)
        layout.setSpacing(8)
        layout.setContentsMargins(18, 18, 18, 18)

        layout.addWidget(QLabel('股票代码'))
        self._code_input = QLineEdit()
        self._code_input.setPlaceholderText('如 000001 / sh600519 / 000001.SZ')
        if prefill_code:
            self._code_input.setText(prefill_code)
        layout.addWidget(self._code_input)

        layout.addWidget(QLabel('名称（可选）'))
        self._name_input = QLineEdit()
        self._name_input.setPlaceholderText('如 贵州茅台')
        if prefill_name:
            self._name_input.setText(prefill_name)
        layout.addWidget(self._name_input)

        _mv = get_default_position_value()
        layout.addWidget(QLabel(f'持有数量（股，留空按{_mv}元市值默认）'))
        self._shares_input = QLineEdit()
        self._shares_input.setPlaceholderText(f'留空=ROUND({_mv}/现价,-2)')
        layout.addWidget(self._shares_input)

        layout.addWidget(QLabel('加入时价（元）'))
        self._cost_input = QLineEdit()
        self._cost_input.setPlaceholderText('留空=加入时现价')
        if prefill_price > 0:
            self._cost_input.setText(_format_price(prefill_price, prefill_code))
        layout.addWidget(self._cost_input)

        layout.addWidget(QLabel('加入分组'))
        self._group_combo = QComboBox()
        self._group_combo.setStyleSheet(_COMBO_STYLE)
        for g in (groups or []):
            self._group_combo.addItem(g)
        self._group_combo.addItem(self._NEW_GROUP_ITEM)
        if self.result_group:
            idx = self._group_combo.findText(self.result_group)
            if idx >= 0:
                self._group_combo.setCurrentIndex(idx)
        self._group_combo.activated.connect(self._maybe_new_group)
        layout.addWidget(self._group_combo)

        layout.addStretch(1)

        btns = QDialogButtonBox(
            QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel
        )
        btns.button(QDialogButtonBox.StandardButton.Ok).setText('确定')
        btns.button(QDialogButtonBox.StandardButton.Cancel).setText('取消')
        btns.accepted.connect(self._on_accept)
        btns.rejected.connect(self.reject)
        layout.addWidget(btns)

    def _on_accept(self):
        raw = self._code_input.text().strip().upper()
        if not raw:
            QMessageBox.warning(self, '提示', '请输入股票代码')
            return

        code = _parse_stock_code(raw)
        if not code:
            QMessageBox.warning(
                self, '提示',
                '请输入有效的6位代码\n（支持 000001 / sh600519 / 000001.SZ 等格式）',
            )
            return

        # 8位期权代码拒绝
        if classify_code(code) != 'stock_or_etf':
            QMessageBox.warning(self, '提示', '期权代码请使用「期权」Tab 添加')
            return

        # 确定市场前缀
        if '.' in raw.upper():
            suffix = raw.upper().split('.', 1)[1]
            if suffix in ('SH', 'SS', 'CSI'):
                code = f'sh{code}'
            elif suffix == 'SZ':
                code = f'sz{code}'
        elif raw.upper().startswith(('SH', 'SZ', 'BJ')):
            prefix = raw[:2].lower()
            code = f'{prefix}{code}'

        sel_group = self._group_combo.currentText()
        if sel_group == self._NEW_GROUP_ITEM:
            sel_group = ''
        for existing in self._existing_by_group.get(sel_group, set()):
            if _code_to_tencent(existing) == _code_to_tencent(code):
                QMessageBox.information(self, '提示', f'代码 {code} 已在该分组中')
                return

        shares_text = self._shares_input.text().strip()
        if shares_text:
            try:
                shares = int(shares_text)
                if shares <= 0:
                    raise ValueError
            except (ValueError, TypeError):
                QMessageBox.warning(self, '提示', '数量格式不正确，请输入正整数，或留空')
                return
        else:
            shares = None

        cost_text = self._cost_input.text().strip()
        if cost_text:
            try:
                cost = float(cost_text)
                if cost <= 0:
                    raise ValueError
            except (ValueError, TypeError):
                QMessageBox.warning(self, '提示', '加入时价格式不正确，请输入正数，或留空')
                return
        else:
            cost = None

        self.result_code = code
        self.result_name = self._name_input.text().strip()
        self.result_shares = shares
        self.result_cost_price = cost
        self.result_theme = ''
        self.result_buy_date = datetime.now().strftime('%Y-%m-%d')
        sel = self._group_combo.currentText()
        if sel and sel != self._NEW_GROUP_ITEM:
            self.result_group = sel
        self.accept()

    def _maybe_new_group(self, _idx: int):
        """选择「新建分组」项 → 弹框输入名称并插入。"""
        if self._group_combo.currentText() != self._NEW_GROUP_ITEM:
            return
        name, ok = QInputDialog.getText(self, '新建分组', '分组名称：')
        name = (name or '').strip()
        new_idx = self._group_combo.findText(self._NEW_GROUP_ITEM)
        if not ok or not name:
            self._group_combo.setCurrentIndex(0)
            return
        exist = self._group_combo.findText(name)
        if exist >= 0:
            self._group_combo.setCurrentIndex(exist)
            return
        self._group_combo.insertItem(new_idx, name)
        self._group_combo.setCurrentIndex(new_idx)


class EditWatchlistDialog(QDialog):
    """编辑自选股：数量、成本价、迁移分组。"""

    def __init__(self, rec: dict, parent=None, groups: list[str] | None = None):
        super().__init__(parent)
        self.setWindowTitle('编辑自选股')
        self.setFixedSize(340, 340)
        self.setStyleSheet(f'background-color: {DARK_BG}; color: #ffffff;')
        self._code = rec['code']
        self.result_shares = None
        self.result_cost_price = None
        self.result_theme = rec.get('theme') or ''
        self.result_group = rec.get('group') or ''

        layout = QVBoxLayout(self)
        layout.setSpacing(8)
        layout.setContentsMargins(18, 18, 18, 18)

        label = f'{_clean_code(rec["code"])}'
        layout.addWidget(QLabel(label))

        layout.addWidget(QLabel('持有数量（股）'))
        self._shares_input = QLineEdit()
        shares = rec.get('shares')
        self._shares_input.setText(str(shares) if shares else '')
        self._shares_input.setPlaceholderText('如 1000')
        layout.addWidget(self._shares_input)

        layout.addWidget(QLabel('加入时价（元）'))
        self._cost_input = QLineEdit()
        cost = rec.get('cost_price') or rec.get('add_price')
        self._cost_input.setText(_format_price(cost, rec.get('code')) if cost else '')
        self._cost_input.setPlaceholderText('留空=加入时现价')
        layout.addWidget(self._cost_input)

        layout.addWidget(QLabel('迁移分组'))
        self._group_combo = QComboBox()
        self._group_combo.setStyleSheet(_COMBO_STYLE)
        for g in (groups or []):
            self._group_combo.addItem(g)
        idx = self._group_combo.findText(self.result_group)
        if idx >= 0:
            self._group_combo.setCurrentIndex(idx)
        layout.addWidget(self._group_combo)

        layout.addStretch(1)

        btns = QDialogButtonBox(
            QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel
        )
        btns.button(QDialogButtonBox.StandardButton.Ok).setText('保存')
        btns.button(QDialogButtonBox.StandardButton.Cancel).setText('取消')
        btns.accepted.connect(self._on_accept)
        btns.rejected.connect(self.reject)
        layout.addWidget(btns)

    def _on_accept(self):
        shares_text = self._shares_input.text().strip()
        if shares_text:
            try:
                shares = int(shares_text)
                if shares <= 0:
                    raise ValueError
            except (ValueError, TypeError):
                QMessageBox.warning(self, '提示', '数量格式不正确，请输入正整数，或留空')
                return
        else:
            shares = None

        cost_text = self._cost_input.text().strip()
        if cost_text:
            try:
                cost = float(cost_text)
                if cost <= 0:
                    raise ValueError
            except (ValueError, TypeError):
                QMessageBox.warning(self, '提示', '加入时价格式不正确，请输入正数，或留空')
                return
        else:
            cost = None

        self.result_shares = shares
        self.result_cost_price = cost
        if self._group_combo.count():
            self.result_group = self._group_combo.currentText()
        self.accept()


class EditTradeDialog(QDialog):
    """编辑自选股：直接修改 / 买入 / 卖出，三种模式。"""

    def __init__(self, rec: dict, parent=None, current_price: float = 0.0, is_position: bool = False):
        super().__init__(parent)
        self.setWindowTitle('编辑自选 / 交易')
        self.setFixedSize(360, 420)
        self.setStyleSheet(f'background-color: {DARK_BG}; color: #ffffff;')
        self._rec = rec
        self._is_position = is_position
        self._cur_price = current_price
        self._code_key = 'stock_code' if is_position else 'code'
        self._name_key = 'stock_name' if is_position else 'name'
        self._result = None

        layout = QVBoxLayout(self)
        layout.setSpacing(8)
        layout.setContentsMargins(18, 18, 18, 18)

        info = QLabel(f'代码: {_clean_code(rec.get(self._code_key, ""))}'
                      f'  名称: {rec.get(self._name_key, "")}')
        info.setStyleSheet('font-weight: bold;')
        layout.addWidget(info)

        self._mode_tabs = QTabWidget()
        self._mode_tabs.setMaximumHeight(32)
        self._mode_tabs.addTab(QWidget(), '直接修改')
        self._mode_tabs.addTab(QWidget(), '买入')
        self._mode_tabs.addTab(QWidget(), '卖出')
        self._mode_tabs.currentChanged.connect(self._on_mode_changed)
        layout.addWidget(self._mode_tabs)

        self._form_stack = QStackedWidget()

        edit_w = QWidget()
        ef = QFormLayout(edit_w)
        ef.setSpacing(6)
        self._shares_edit = QSpinBox()
        self._shares_edit.setRange(0, 999999999)
        cur_shares = int(rec.get('shares', 0) or 0)
        self._shares_edit.setValue(cur_shares)
        ef.addRow('数量:', self._shares_edit)
        self._cost_edit = QDoubleSpinBox()
        self._cost_edit.setRange(0.0, 999999.99)
        self._cost_edit.setDecimals(3)
        self._cost_edit.setValue(float(rec.get('cost_price', 0) or 0))
        ef.addRow('成本价:', self._cost_edit)
        if not is_position:
            self._theme_edit = QLineEdit(rec.get('theme', ''))
            ef.addRow('板块:', self._theme_edit)
        else:
            self._theme_edit = None
        self._form_stack.addWidget(edit_w)

        trade_w = QWidget()
        tf = QFormLayout(trade_w)
        tf.setSpacing(6)
        self._trade_price = QDoubleSpinBox()
        self._trade_price.setRange(0.0, 999999.99)
        self._trade_price.setDecimals(2)
        self._trade_price.setValue(current_price if current_price > 0
                                   else float(rec.get('cost_price', 0) or 0))
        tf.addRow('成交价:', self._trade_price)
        self._trade_shares = QSpinBox()
        self._trade_shares.setRange(100, 999999999)
        self._trade_shares.setSingleStep(100)
        self._trade_shares.setValue(100)
        tf.addRow('交易数量:', self._trade_shares)
        self._trade_date = QDateEdit(calendarPopup=True)
        self._trade_date.setDate(QDate.currentDate())
        self._trade_date.setDisplayFormat('yyyy-MM-dd')
        tf.addRow('交易日期:', self._trade_date)
        self._form_stack.addWidget(trade_w)

        layout.addWidget(self._form_stack)

        self._preview = QLabel()
        self._preview.setStyleSheet('background:#1a1a30; padding:8px; font-size:10pt; border-radius:4px;')
        layout.addWidget(self._preview)

        btns = QDialogButtonBox(
            QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel
        )
        btns.button(QDialogButtonBox.StandardButton.Ok).setText('确认')
        btns.button(QDialogButtonBox.StandardButton.Cancel).setText('取消')
        btns.accepted.connect(self._on_ok)
        btns.rejected.connect(self.reject)
        layout.addWidget(btns)

        self._mode_tabs.setCurrentIndex(0)
        self._update_preview()

    def _on_mode_changed(self, idx):
        if idx == 0:
            self._form_stack.setCurrentIndex(0)
        else:
            self._form_stack.setCurrentIndex(1)
        self._update_preview()

    def _update_preview(self):
        mode = self._mode_tabs.currentIndex()
        cur_shares = int(self._rec.get('shares', 0) or 0)
        cur_cost = float(self._rec.get('cost_price', 0) or 0)

        if mode == 0:
            ns = self._shares_edit.value()
            nc = self._cost_edit.value()
            self._preview.setText(f'修改后: 数量 {ns}股, 成本 ¥{_format_price(nc, self._rec.get("code"))}')
        elif mode == 1:
            price = self._trade_price.value()
            shares = self._trade_shares.value()
            new_s, new_c = compute_buy(cur_shares, cur_cost, shares, price)
            self._preview.setText(f'买入后: 数量 {new_s}股, 加权均价 ¥{_format_price(new_c, self._rec.get("code"))}')
        else:
            shares = self._trade_shares.value()
            new_s = compute_sell(cur_shares, shares)
            if new_s is None:
                self._preview.setStyleSheet('background:#3a1010; padding:8px; font-size:10pt; border-radius:4px;')
                self._preview.setText(f'错误: 卖出 {shares}股 超过当前持仓 {cur_shares}股')
                return
            self._preview.setStyleSheet('background:#1a1a30; padding:8px; font-size:10pt; border-radius:4px;')
            self._preview.setText(f'卖出后: 剩余 {new_s}股, 成本不变 ¥{_format_price(cur_cost, self._rec.get("code"))}')

    def _on_ok(self):
        mode = self._mode_tabs.currentIndex()
        cur_shares = int(self._rec.get('shares', 0) or 0)
        cur_cost = float(self._rec.get('cost_price', 0) or 0)

        if mode == 0:
            self._result = {
                'mode': 'edit',
                'new_shares': self._shares_edit.value(),
                'new_cost': self._cost_edit.value(),
                'theme': self._theme_edit.text() if self._theme_edit else self._rec.get('theme', ''),
                'trade_record': None,
            }
        elif mode == 1:
            price = self._trade_price.value()
            shares = self._trade_shares.value()
            new_s, new_c = compute_buy(cur_shares, cur_cost, shares, price)
            self._result = {
                'mode': 'buy',
                'new_shares': new_s,
                'new_cost': new_c,
                'theme': self._rec.get('theme', ''),
                'trade_record': {
                    'date': self._trade_date.date().toString('yyyy-MM-dd'),
                    'direction': 'buy', 'shares': shares,
                    'price': price, 'timestamp': datetime.now().isoformat(),
                },
            }
        else:
            price = self._trade_price.value()
            shares = self._trade_shares.value()
            new_s = compute_sell(cur_shares, shares)
            if new_s is None:
                QMessageBox.warning(self, '错误', f'卖出 {shares}股 超过当前持仓 {cur_shares}股')
                return
            self._result = {
                'mode': 'sell',
                'new_shares': new_s,
                'new_cost': cur_cost,
                'theme': self._rec.get('theme', ''),
                'trade_record': {
                    'date': self._trade_date.date().toString('yyyy-MM-dd'),
                    'direction': 'sell', 'shares': shares,
                    'price': price, 'timestamp': datetime.now().isoformat(),
                },
            }
        self.accept()

    def get_result(self) -> dict | None:
        return self._result


class _OptionPositionDialog(QDialog):
    """期权持仓台账 — 添加/编辑弹窗。"""

    def __init__(self, parent=None, rec: dict | None = None):
        super().__init__(parent)
        self._rec = rec
        is_edit = rec is not None
        self.setWindowTitle('编辑期权持仓' if is_edit else '添加期权持仓')
        self.setMinimumWidth(440)
        self.setStyleSheet(f'background-color: {DARK_BG}; color: #ffffff;')

        layout = QVBoxLayout(self)
        layout.setSpacing(8)
        layout.setContentsMargins(18, 16, 18, 16)

        top_hint = QLabel(
            '期权价格通常按每份报价，实际权利金 = 价格 × 合约单位 × 张数。'
            '昨收按交易所前结价计算。'
        )
        top_hint.setStyleSheet('color: #f0c040; font-size: 11px; padding: 4px;')
        top_hint.setWordWrap(True)
        layout.addWidget(top_hint)

        self._short_warning = QLabel(
            '卖出义务仓存在保证金、被指派和理论无限风险。'
            '本系统第一版仅记录，不提供保证金和风险测算。'
        )
        self._short_warning.setStyleSheet(
            'color: #ff6b6b; background-color: #3a1010; font-size: 11px; '
            'padding: 6px; border-radius: 4px; font-weight: bold;'
        )
        self._short_warning.setWordWrap(True)
        self._short_warning.setVisible(False)
        layout.addWidget(self._short_warning)

        form = QFormLayout()
        form.setSpacing(8)

        self._code_edit = QLineEdit(rec.get('code', '') if rec else '')
        self._code_edit.setPlaceholderText('8位代码或完整合约简称…')
        self._code_edit.setStyleSheet(_LINE_EDIT_STYLE)
        form.addRow('合约代码:', self._code_edit)

        self._name_edit = QLineEdit(rec.get('name', '') if rec else '')
        self._name_edit.setPlaceholderText('合约简称或备注…')
        self._name_edit.setStyleSheet(_LINE_EDIT_STYLE)
        form.addRow('合约简称/备注:', self._name_edit)

        self._underlying_code_edit = QLineEdit(rec.get('underlying_code', '') if rec else '')
        self._underlying_code_edit.setPlaceholderText('如 510300')
        self._underlying_code_edit.setStyleSheet(_LINE_EDIT_STYLE)
        form.addRow('标的代码:', self._underlying_code_edit)

        self._underlying_name_edit = QLineEdit(rec.get('underlying_name', '') if rec else '')
        self._underlying_name_edit.setPlaceholderText('如 沪深300ETF')
        self._underlying_name_edit.setStyleSheet(_LINE_EDIT_STYLE)
        form.addRow('标的名称:', self._underlying_name_edit)

        self._option_type_combo = QComboBox()
        self._option_type_combo.addItems(['认购 Call', '认沽 Put'])
        if rec and rec.get('option_type') == 'put':
            self._option_type_combo.setCurrentIndex(1)
        self._option_type_combo.setStyleSheet(_COMBO_STYLE)
        form.addRow('方向类型:', self._option_type_combo)

        self._position_side_combo = QComboBox()
        self._position_side_combo.addItems(['买入权利仓', '卖出义务仓'])
        if rec and rec.get('position_side') == 'short_obligation':
            self._position_side_combo.setCurrentIndex(1)
        self._position_side_combo.setStyleSheet(_COMBO_STYLE)
        self._position_side_combo.currentIndexChanged.connect(self._on_side_changed)
        form.addRow('持仓方向:', self._position_side_combo)

        self._strike_spin = QDoubleSpinBox()
        self._strike_spin.setRange(0.0, 99999.99)
        self._strike_spin.setDecimals(4)
        self._strike_spin.setValue(float(rec.get('strike_price', 0) or 0) if rec else 0.0)
        _apply_numeric_spin_style(self._strike_spin)
        form.addRow('行权价:', self._strike_spin)

        self._expiry_date_edit = QDateEdit(calendarPopup=True)
        self._expiry_date_edit.setDisplayFormat('yyyy-MM-dd')
        if rec and rec.get('expiry_date'):
            try:
                d = QDate.fromString(rec['expiry_date'][:10], 'yyyy-MM-dd')
                if d.isValid():
                    self._expiry_date_edit.setDate(d)
            except Exception:
                self._expiry_date_edit.setDate(QDate.currentDate())
        else:
            self._expiry_date_edit.setDate(QDate.currentDate())
        self._expiry_date_edit.setStyleSheet(_SPIN_STYLE)
        form.addRow('到期日:', self._expiry_date_edit)

        self._contract_unit_spin = QSpinBox()
        self._contract_unit_spin.setRange(1, 999999)
        self._contract_unit_spin.setValue(
            int(rec.get('contract_unit', 10000) or 10000) if rec else 10000
        )
        _apply_numeric_spin_style(self._contract_unit_spin)
        form.addRow('合约单位:', self._contract_unit_spin)

        self._contracts_spin = QSpinBox()
        self._contracts_spin.setRange(1, 999999)
        self._contracts_spin.setValue(int(rec.get('contracts', 1) or 1) if rec else 1)
        _apply_numeric_spin_style(self._contracts_spin)
        form.addRow('张数:', self._contracts_spin)

        self._open_price_spin = QDoubleSpinBox()
        self._open_price_spin.setRange(0.0, 999999.99)
        self._open_price_spin.setDecimals(6)
        self._open_price_spin.setValue(float(rec.get('open_price', 0) or 0) if rec else 0.0)
        _apply_numeric_spin_style(self._open_price_spin)
        form.addRow('开仓价格:', self._open_price_spin)

        self._current_price_spin = QDoubleSpinBox()
        self._current_price_spin.setRange(0.0, 999999.99)
        self._current_price_spin.setDecimals(6)
        self._current_price_spin.setValue(float(rec.get('current_price', 0) or 0) if rec else 0.0)
        _apply_numeric_spin_style(self._current_price_spin)
        self._current_price_spin.setSpecialValueText('未维护')
        form.addRow('当前价格:', self._current_price_spin)

        self._pre_close_spin = QDoubleSpinBox()
        self._pre_close_spin.setRange(0.0, 999999.99)
        self._pre_close_spin.setDecimals(6)
        self._pre_close_spin.setValue(float(rec.get('pre_close', 0) or 0) if rec else 0.0)
        _apply_numeric_spin_style(self._pre_close_spin)
        self._pre_close_spin.setSpecialValueText('未维护')
        form.addRow('昨收/前结价:', self._pre_close_spin)

        self._open_date_edit = QDateEdit(calendarPopup=True)
        self._open_date_edit.setDisplayFormat('yyyy-MM-dd')
        if rec and rec.get('open_date'):
            try:
                d = QDate.fromString(rec['open_date'][:10], 'yyyy-MM-dd')
                if d.isValid():
                    self._open_date_edit.setDate(d)
            except Exception:
                self._open_date_edit.setDate(QDate.currentDate())
        else:
            self._open_date_edit.setDate(QDate.currentDate())
        self._open_date_edit.setStyleSheet(_SPIN_STYLE)
        form.addRow('开仓日期:', self._open_date_edit)

        self._status_combo = QComboBox()
        self._status_combo.addItems(['持有中', '已平仓', '已到期', '已行权', '已作废'])
        if rec:
            status_map = {'holding': 0, 'closed': 1, 'expired': 2, 'exercised': 3, 'void': 4}
            self._status_combo.setCurrentIndex(status_map.get(rec.get('status', 'holding'), 0))
        self._status_combo.setStyleSheet(_COMBO_STYLE)
        form.addRow('状态:', self._status_combo)

        self._risk_note_edit = QLineEdit(rec.get('risk_note', '') if rec else '')
        self._risk_note_edit.setPlaceholderText('风险备注（可选）…')
        self._risk_note_edit.setStyleSheet(_LINE_EDIT_STYLE)
        form.addRow('风险备注:', self._risk_note_edit)

        layout.addLayout(form)

        btns = QDialogButtonBox(
            QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel
        )
        btns.button(QDialogButtonBox.StandardButton.Ok).setText('保存')
        btns.button(QDialogButtonBox.StandardButton.Cancel).setText('取消')
        btns.accepted.connect(self._on_accept)
        btns.rejected.connect(self.reject)
        layout.addWidget(btns)

        self._on_side_changed(self._position_side_combo.currentIndex())

    def _on_side_changed(self, idx):
        self._short_warning.setVisible(idx == 1)

    def _on_accept(self):
        code = self._code_edit.text().strip()
        if not code:
            QMessageBox.warning(self, '缺少合约代码', '合约代码不能为空。')
            return
        if code.isdigit() and len(code) == 6:
            QMessageBox.warning(
                self, '代码不匹配',
                '这是股票/ETF代码，请到自选股添加；期权合约请填写8位代码或完整合约简称。',
            )
            return
        if self._contract_unit_spin.value() <= 0:
            QMessageBox.warning(self, '合约单位', '合约单位必须大于 0。')
            return
        if self._contracts_spin.value() <= 0:
            QMessageBox.warning(self, '张数', '张数必须大于 0。')
            return
        self.accept()

    def get_record(self) -> dict:
        status_map = {0: 'holding', 1: 'closed', 2: 'expired', 3: 'exercised', 4: 'void'}
        now = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        rec = {
            'code': self._code_edit.text().strip(),
            'name': self._name_edit.text().strip(),
            'underlying_code': self._underlying_code_edit.text().strip(),
            'underlying_name': self._underlying_name_edit.text().strip(),
            'option_type': 'put' if self._option_type_combo.currentIndex() == 1 else 'call',
            'position_side': 'short_obligation' if self._position_side_combo.currentIndex() == 1 else 'long_right',
            'strike_price': self._strike_spin.value(),
            'expiry_date': self._expiry_date_edit.date().toString('yyyy-MM-dd'),
            'contract_unit': self._contract_unit_spin.value(),
            'contracts': self._contracts_spin.value(),
            'open_price': self._open_price_spin.value(),
            'current_price': self._current_price_spin.value(),
            'pre_close': self._pre_close_spin.value(),
            'open_date': self._open_date_edit.date().toString('yyyy-MM-dd'),
            'status': status_map[self._status_combo.currentIndex()],
            'risk_note': self._risk_note_edit.text().strip(),
            'updated_at': now,
        }
        if self._rec:
            rec.setdefault('created_at', self._rec.get('created_at', now))
        else:
            rec['created_at'] = now
        return rec


class _QuickHoldingDialog(QDialog):
    """从自选股 / 持仓视角一键加入真实持仓，预填代码、名称、参考价。"""

    def __init__(self, code: str, name: str, ref_price: float, parent=None,
                 ref_shares: int | None = None):
        super().__init__(parent)
        self.setWindowTitle('加入真实持仓')
        self.setFixedSize(340, 380)
        self.setStyleSheet(f'background-color: {DARK_BG}; color: #ffffff;')
        self._result: dict | None = None
        self._code = code
        self._ref_price = ref_price

        layout = QVBoxLayout(self)
        layout.setSpacing(8)
        layout.setContentsMargins(18, 18, 18, 18)

        layout.addWidget(QLabel('股票代码'))
        self._code_input = QLineEdit(code)
        self._code_input.setReadOnly(True)
        self._code_input.setStyleSheet('color: #888888;')
        layout.addWidget(self._code_input)

        layout.addWidget(QLabel('名称（可选）'))
        self._name_input = QLineEdit(name or '')
        self._name_input.setPlaceholderText('如 贵州茅台')
        layout.addWidget(self._name_input)

        _mv = get_default_position_value()
        layout.addWidget(QLabel(f'持有数量（股，留空按{_mv}元市值默认）'))
        self._shares_input = QLineEdit()
        self._shares_input.setPlaceholderText(f'留空=ROUND({_mv}/现价,-2)')
        if ref_shares and ref_shares > 0:
            self._shares_input.setText(str(int(ref_shares)))
        layout.addWidget(self._shares_input)

        layout.addWidget(QLabel('加入时价（元）'))
        self._cost_input = QLineEdit()
        self._cost_input.setPlaceholderText('留空=加入时现价')
        if ref_price > 0:
            self._cost_input.setText(_format_price(ref_price, code))
        layout.addWidget(self._cost_input)

        layout.addStretch(1)

        btns = QDialogButtonBox(
            QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel
        )
        btns.button(QDialogButtonBox.StandardButton.Ok).setText('加入持仓')
        btns.button(QDialogButtonBox.StandardButton.Cancel).setText('取消')
        btns.accepted.connect(self._on_accept)
        btns.rejected.connect(self.reject)
        layout.addWidget(btns)

    def _on_accept(self):
        cost_text = self._cost_input.text().strip()
        if cost_text:
            try:
                cost = float(cost_text)
                if cost <= 0:
                    raise ValueError
            except (ValueError, TypeError):
                QMessageBox.warning(self, '提示', '加入时价格式不正确，请输入正数，或留空')
                return
        else:
            cost = self._ref_price
        if cost <= 0:
            QMessageBox.warning(self, '提示', '无法获取现价，请手动填写加入时价')
            return

        shares_text = self._shares_input.text().strip()
        if shares_text:
            try:
                shares = int(shares_text)
                if shares <= 0:
                    raise ValueError
            except (ValueError, TypeError):
                QMessageBox.warning(self, '提示', '数量格式不正确，请输入正整数，或留空')
                return
        else:
            shares = _default_wl_shares(cost)
            if not shares:
                QMessageBox.warning(self, '提示', '无法按现价估算数量，请手动填写持有数量')
                return

        self._result = {
            'code': self._code,
            'name': self._name_input.text().strip(),
            'cost_price': cost,
            'shares': shares,
        }
        self.accept()

    def get_data(self) -> dict | None:
        return self._result


# ── PnL 日历组件 ──

class _WatchlistPnlCalendarWidget(QWidget):
    """自选股收益月历：顶部汇总 + 月份导航 + 日历格子。"""

    def __init__(self, parent=None):
        super().__init__(parent)
        today = date.today()
        self._current_year = today.year
        self._current_month = today.month
        self._records = {}
        self._setup_ui()

    def _setup_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(2, 2, 2, 2)
        layout.setSpacing(1)

        self._summary_label = QLabel('本月盈亏: --  今日盈亏: --  总市值: --  总成本: --  缺失: 0只')
        self._summary_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self._summary_label.setStyleSheet('font-size:9pt; font-weight:bold; padding:2px;')
        layout.addWidget(self._summary_label)

        nav_style = (
            'QPushButton { background:#1a2a3a; color:#ccd; border:1px solid #2a3a4a;'
            ' border-radius:3px; font-size:10pt; padding:1px 4px; }'
            'QPushButton:hover { background:#2a3a5a; color:#fff; }'
        )

        nav = QHBoxLayout()
        self._prev_btn = QPushButton('◀')
        self._prev_btn.setFixedSize(28, 22)
        self._prev_btn.setStyleSheet(nav_style)
        self._prev_btn.clicked.connect(self._prev_month)
        self._month_label = QLabel(f'{self._current_year}年{self._current_month}月')
        self._month_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self._month_label.setStyleSheet('font-size:10pt; font-weight:bold;')
        self._today_btn = QPushButton('本月')
        self._today_btn.setFixedSize(36, 22)
        self._today_btn.setStyleSheet(nav_style)
        self._today_btn.clicked.connect(self._go_today)
        self._next_btn = QPushButton('▶')
        self._next_btn.setFixedSize(28, 22)
        self._next_btn.setStyleSheet(nav_style)
        self._next_btn.clicked.connect(self._next_month)
        nav.addWidget(self._prev_btn)
        nav.addWidget(self._month_label, 1)
        nav.addWidget(self._today_btn)
        nav.addWidget(self._next_btn)
        layout.addLayout(nav)

        dow = QHBoxLayout()
        for day in ['一', '二', '三', '四', '五', '六', '日']:
            lbl = QLabel(day)
            lbl.setAlignment(Qt.AlignmentFlag.AlignCenter)
            lbl.setStyleSheet('font-weight:bold; color:#778; font-size:8pt;')
            dow.addWidget(lbl)
        layout.addLayout(dow)

        self._grid = QGridLayout()
        self._grid.setSpacing(1)
        layout.addLayout(self._grid)
        self._rebuild_grid()

    def _prev_month(self):
        if self._current_month == 1:
            self._current_month = 12
            self._current_year -= 1
        else:
            self._current_month -= 1
        self._refresh()

    def _next_month(self):
        if self._current_month == 12:
            self._current_month = 1
            self._current_year += 1
        else:
            self._current_month += 1
        self._refresh()

    def _go_today(self):
        today = date.today()
        self._current_year = today.year
        self._current_month = today.month
        self._refresh()

    def set_records(self, records: list[dict]):
        self._records = {r['date']: r for r in records}
        self._refresh()

    def summary_text(self) -> str:
        return self._summary_label.text()

    def _refresh(self):
        self._month_label.setText(f'{self._current_year}年{self._current_month}月')
        self._rebuild_grid()
        self._update_summary()

    def _rebuild_grid(self):
        while self._grid.count():
            item = self._grid.takeAt(0)
            if item.widget():
                item.widget().deleteLater()

        import calendar
        first_day = date(self._current_year, self._current_month, 1)
        first_weekday = first_day.weekday()
        days_in_month = calendar.monthrange(self._current_year, self._current_month)[1]

        row, col = 0, first_weekday
        for d in range(1, days_in_month + 1):
            date_str = f'{self._current_year}-{self._current_month:02d}-{d:02d}'
            rec = self._records.get(date_str, {})
            pnl = rec.get('total_pnl', 0) or 0
            count = rec.get('stock_count', 0) or 0

            cell = QLabel()
            cell.setAlignment(Qt.AlignmentFlag.AlignCenter)
            cell.setMinimumSize(30, 28)

            if count > 0:
                text = f'{d}\n{pnl:+.0f}\n{count}只'
                color = '#e84444' if pnl > 0 else ('#2ecc71' if pnl < 0 else '#aaaaaa')
                tip = f'{date_str}\n盈亏: {pnl:+,.0f}\n涉及: {count}只'
            else:
                text = f'{d}\n—'
                color = '#555555'
                tip = f'{date_str}\n无数据'
            cell.setText(text)
            cell.setStyleSheet(
                f'border:1px solid #2a2a4a; padding:0px; color:{color}; font-size:8pt;'
            )
            cell.setToolTip(tip)
            self._grid.addWidget(cell, row, col)
            col += 1
            if col > 6:
                col = 0
                row += 1

    def _update_summary(self):
        prefix = f'{self._current_year}-{self._current_month:02d}'
        month_recs = [r for d, r in self._records.items() if d.startswith(prefix)]
        total_pnl = sum(r.get('total_pnl', 0) or 0 for r in month_recs)
        today_str = date.today().strftime('%Y-%m-%d')
        today_rec = self._records.get(today_str, {})
        today_pnl = today_rec.get('total_pnl', 0) or 0
        missing = today_rec.get('missing_count', 0)
        total_val = 0.0
        total_cst = 0.0
        for r in month_recs:
            s = r.get('summary', {})
            total_val += s.get('total_value', 0) or 0
            total_cst += s.get('total_cost', 0) or 0
        self._summary_label.setText(
            f'本月盈亏: {total_pnl:+,.0f}  今日盈亏: {today_pnl:+,.0f}  '
            f'总市值: {self._fmt(total_val)}  总成本: {self._fmt(total_cst)}  缺失: {missing}只'
        )

    @staticmethod
    def _fmt(v: float) -> str:
        if abs(v) >= 1e8:
            return f'{v/1e8:.1f}亿'
        if abs(v) >= 1e4:
            return f'{v/1e4:.0f}万'
        return f'{v:.0f}'


class _EtfFileUpdateWorker(QThread):
    """后台就地刷新所选 xlsm 的实时数据格，避免大文件卡 UI。"""
    finished = Signal(dict)

    def __init__(self, paths: list[str], parent=None):
        super().__init__(parent)
        self._paths = paths

    def run(self):
        try:
            from core.etf_file_updater import update_etf_files
            res = update_etf_files(self._paths)
        except Exception as e:
            res = {'results': [{'path': p, 'ok': False, 'msg': str(e)} for p in self._paths],
                   'quotes': 0}
        self.finished.emit(res)


class StockPanel(QWidget):
    status_changed = Signal(str, str)
    point_count_changed = Signal(int)
    bottom_status_changed = Signal(str)

    def __init__(self, parent=None):
        super().__init__(parent)
        self.worker = None
        self.worker_out = None
        self.worker_gain = None
        self.worker_drop = None
        self.worker_watchlist = None
        self._pending_period = None
        self.period = '今日'
        self._last_in_snap = None
        self._last_out_snap = None
        self._last_gain_snap = None
        self._last_drop_snap = None
        self._watchlist_codes = load_watchlist()
        self._wl_default_group = '新建分组1'
        self._ensure_watchlist_groups()
        self._wl_active_group: str | None = None  # None = 全部
        self._last_watchlist_snap = None
        # 自选股表格排序状态：涨跌幅(col 6) / 主力净额(col 7) / 累计盈亏(col 11) / 盈亏比例(col 12) 可点击排序
        # None 表示未排序；首次点击 → 降序；同列再点 → 升序（降↔升切换）
        self._wl_sort_col: int | None = None
        self._wl_sort_desc: bool = True
        self._last_option_refresh = 0.0

        root = QVBoxLayout(self)
        root.setContentsMargins(18, 16, 18, 16)
        root.setSpacing(12)

        top = QHBoxLayout()
        top.setSpacing(10)

        title = QLabel('📈 个股资金')
        title.setStyleSheet('color: #ffffff; font-size: 18px; font-weight: bold;')
        top.addWidget(title)
        top.addStretch(1)

        label = QLabel('周期')
        label.setStyleSheet(f'color: {MUTED}; font-size: 12px;')
        top.addWidget(label)

        self.period_combo = QComboBox()
        self.period_combo.addItems(list(_PERIODS.keys()))
        self.period_combo.setFixedWidth(90)
        self.period_combo.currentTextChanged.connect(self._on_period_changed)
        top.addWidget(self.period_combo)

        # 更新 ETF 数据库：配置 TickFlow key 或演示码时显示。
        # 一键优先用 TickFlow（演示码/失败时走免费兜底）把当天行情就地写进所选 xlsm。
        self._etf_update_btn = QPushButton('📊 更新 ETF 数据库', self)
        self._etf_update_btn.setFixedHeight(28)
        self._etf_update_btn.setToolTip(
            '优先用 TickFlow 行情就地刷新所选 .xlsm 的实时数据单元格；演示码/失败时走免费兜底\n'
            '（保留宏、按钮、公式与其它数据，可一次多选两个文件）'
        )
        self._etf_update_btn.setStyleSheet(
            'QPushButton { background-color: #1a3a2a; color: #7aebb8;'
            ' border: 1px solid #3a8a6a; border-radius: 4px; padding: 0 12px; font-size: 12px; }'
            'QPushButton:hover { background-color: #2a6a4a; color: #fff; }'
        )
        self._etf_update_btn.clicked.connect(self._on_update_etf_files)
        top.addWidget(self._etf_update_btn)
        self._refresh_etf_update_btn_visible()

        root.addLayout(top)

        self.tabs = QTabWidget()
        self.tabs.setStyleSheet(
            "QTabWidget::pane { border: none; background-color: " + CHART_BG + "; }"
            "QTabBar::tab { background: #1a1a2e; color: #aaa; padding: 6px 18px; border: 1px solid #333; }"
            "QTabBar::tab:selected { background: #0f3460; color: #fff; border-bottom: 2px solid #e94560; }"
        )
        self.table_in = self._make_stock_table()
        self.table_out = self._make_stock_table()
        self.table_gain = self._make_stock_table()
        self.table_drop = self._make_stock_table()
        self.tabs.addTab(self.table_in, '净流入 TOP30')
        self.tabs.addTab(self.table_out, '净流出 TOP30')
        self.tabs.addTab(self.table_gain, '涨幅 TOP30')
        self.tabs.addTab(self.table_drop, '跌幅 TOP30')
        self.watchlist_widget = self._make_watchlist_tab()
        self.tabs.addTab(self.watchlist_widget, '自选股')
        from ui.ipo_panel import IpoPanel as _IpoPanel
        try:
            self._ipo_widget = _IpoPanel()
            self.tabs.addTab(self._ipo_widget, '🆕 新股')
        except Exception:
            import traceback
            traceback.print_exc()
            self._ipo_widget = QLabel('新股面板加载失败')
        try:
            self.option_widget = self._setup_option_tab()
        except Exception:
            import traceback
            traceback.print_exc()
            self.option_widget = QLabel('期权面板加载失败')
        self.tabs.addTab(self.option_widget, '期权')
        self.tabs.currentChanged.connect(self._on_tab_changed)
        root.addWidget(self.tabs, stretch=1)

        self.meta_lbl = QLabel('等待首次刷新…')
        self.meta_lbl.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.meta_lbl.setStyleSheet(f'color: {MUTED}; font-size: 16px; font-weight: bold;')
        root.addWidget(self.meta_lbl)

        self.setStyleSheet(f'background-color: {DARK_BG};')
        try:
            self._load_cached_period(self.period)
        except Exception:
            import traceback
            traceback.print_exc()
        try:
            self._load_cached_watchlist()
        except Exception:
            import traceback
            traceback.print_exc()

        # TickFlow key 或演示码启用时，自选股 ETF 走 5 秒实时刷新门控。
        self._etf_fast_timer = QTimer(self)
        self._etf_fast_timer.setInterval(5000)
        self._etf_fast_timer.timeout.connect(self._etf_fast_tick)
        self._etf_fast_timer.start()

        self._option_auto_timer = QTimer(self)
        self._option_auto_timer.setInterval(OPTION_AUTO_REFRESH_INTERVAL_MS)
        self._option_auto_timer.timeout.connect(self._option_auto_refresh_tick)
        self._option_auto_timer.start()

    def _refresh_etf_update_btn_visible(self):
        """无 TickFlow key / 演示码不显示「更新 ETF 数据库」按钮。"""
        try:
            from core.credentials import has_tickflow_token
            self._etf_update_btn.setVisible(has_tickflow_token())
        except Exception:
            self._etf_update_btn.setVisible(False)

    def _on_update_etf_files(self):
        """多选待更新的 xlsm，后台优先用 TickFlow、失败走免费兜底刷新实时数据格。"""
        paths, _ = QFileDialog.getOpenFileNames(
            self, '选择要更新的 ETF 数据文件（可多选）',
            '', 'Excel 启用宏工作簿 (*.xlsm)',
        )
        if not paths:
            return
        self._etf_update_btn.setEnabled(False)
        self._etf_update_btn.setText('⏳ 正在更新…')
        self._etf_upd_worker = _EtfFileUpdateWorker(paths)
        self._etf_upd_worker.finished.connect(self._on_etf_files_updated)
        self._etf_upd_worker.start()

    def _on_etf_files_updated(self, res: dict):
        self._etf_update_btn.setEnabled(True)
        self._etf_update_btn.setText('📊 更新 ETF 数据库')
        lines = []
        for r in res.get('results', []):
            name = os.path.basename(r.get('path', ''))
            if r.get('ok'):
                lines.append(f'✅ {name}：更新 {r.get("updated", 0)} 行，跳过 {r.get("skipped", 0)} 行')
            else:
                lines.append(f'❌ {name}：{r.get("msg", "失败")}')
        body = '\n'.join(lines) if lines else '没有可更新的文件'
        body += f'\n\n本次取得行情 {res.get("quotes", 0)} 条。'
        any_ok = any(r.get('ok') for r in res.get('results', []))
        box = QMessageBox(self)
        box.setWindowTitle('ETF 数据库更新结果')
        box.setIcon(QMessageBox.Icon.Information if any_ok else QMessageBox.Icon.Warning)
        box.setText(body)
        box.exec()

    def _etf_fast_tick(self):
        """5 秒定时器：仅在配置了 TickFlow key / 演示码且交易时段时刷新自选股行情。"""
        try:
            from core.credentials import has_tickflow_token
            if not has_tickflow_token() or not is_trading_time():
                return
            if not any(is_etf_code(r.get('code', '')) for r in self._watchlist_codes):
                return
            self._fetch_watchlist()
        except Exception:
            pass

    def _is_option_tab_active(self):
        return (
            self.isVisible()
            and getattr(self, 'tabs', None) is not None
            and self.tabs.currentWidget() is getattr(self, 'option_widget', None)
        )

    def _option_auto_refresh_tick(self):
        if not is_option_auto_refresh_time():
            return
        self._option_refresh_quotes(silent=True, force=True)

    def _refresh_option_on_entry(self):
        if self._is_option_tab_active():
            self._option_refresh_quotes(silent=True, force=True)

    def showEvent(self, event):
        super().showEvent(event)
        if getattr(self, 'tabs', None) is not None:
            QTimer.singleShot(0, self._refresh_option_on_entry)

    def refresh(self):
        """手动刷新（任何时间均可触发，启动时也调用）"""
        self._refresh_etf_update_btn_visible()
        self._fetch()
        self._fetch_watchlist()
        self._ipo_widget.refresh()
        if self._is_option_tab_active():
            QTimer.singleShot(0, self._refresh_option_on_entry)

    def auto_refresh(self):
        if is_trading_time():
            self._fetch()
            self._fetch_watchlist()
            self._ipo_widget.auto_refresh()

    def background_auto_refresh(self):
        """Refresh background-safe quote surfaces without pulling TOP lists."""
        if is_trading_time():
            self._fetch_watchlist()

    def _on_tab_changed(self, index):
        """Tab 切换到「新股」「期权」时触发一次刷新。"""
        w = self.tabs.widget(index)
        if w is self._ipo_widget:
            self._ipo_widget.refresh()
        elif w is self.option_widget:
            QTimer.singleShot(0, self._refresh_option_on_entry)

    def current_point_count(self):
        return self.table_in.rowCount()

    def _make_stock_table(self):
        table = QTableWidget(0, len(_HEADERS))
        table.setHorizontalHeaderLabels(_HEADERS)
        configure_no_truncation(table)
        for i, hdr in enumerate(_HEADERS):
            tip = _HEADER_TIPS.get(hdr)
            if tip:
                table.horizontalHeaderItem(i).setToolTip(tip)
        table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        table.setSelectionMode(QAbstractItemView.SelectionMode.SingleSelection)
        table.verticalHeader().setVisible(False)
        table.setAlternatingRowColors(True)
        table.setShowGrid(False)
        table.setStyleSheet(f'''
            QTableWidget {{
                background-color: {CHART_BG};
                alternate-background-color: #1d2d4a;
                color: #dddddd;
                border: none;
                border-radius: 8px;
                font-size: 12px;
                selection-background-color: #2a3a5e;
            }}
            QHeaderView::section {{
                background-color: #111a30;
                color: #aaaaaa;
                border: none;
                padding: 8px 6px;
                font-weight: bold;
            }}
            QTableWidget::item {{
                padding: 6px 4px;
                border: none;
            }}
        ''')
        h = table.horizontalHeader()
        h.setSectionResizeMode(QHeaderView.ResizeMode.ResizeToContents)
        h.setSectionResizeMode(2, QHeaderView.ResizeMode.Stretch)
        last_col = len(_HEADERS) - 1
        h.setSectionResizeMode(last_col, QHeaderView.ResizeMode.Fixed)
        table.setColumnWidth(last_col, 72)
        table.setContextMenuPolicy(Qt.ContextMenuPolicy.CustomContextMenu)
        table.customContextMenuRequested.connect(lambda pos, t=table: self._on_flow_context_menu(pos, t))
        return table

    def _on_period_changed(self, text):
        self.period = text
        self._fetch()

    def _cache_name(self, direction, period=None):
        return f"stock_{direction}_{period or self.period}"

    def _load_cached_period(self, period):
        in_snap = _load_latest_stock_history_snapshot('in', period)
        out_snap = _load_latest_stock_history_snapshot('out', period)
        loaded = False
        try:
            if in_snap and in_snap.get('rows'):
                self._last_in_snap = in_snap
                self._fill_table(self.table_in, in_snap['rows'])
                loaded = True
            if out_snap and out_snap.get('rows'):
                self._last_out_snap = out_snap
                self._fill_table(self.table_out, out_snap['rows'])
                loaded = True
            if self._last_gain_snap is None:
                gain_snap = load_panel_cache('stock_gain')
                if gain_snap and gain_snap.get('rows'):
                    self._last_gain_snap = gain_snap
                    self._fill_table(self.table_gain, gain_snap['rows'])
            if self._last_drop_snap is None:
                drop_snap = load_panel_cache('stock_drop')
                if drop_snap and drop_snap.get('rows'):
                    self._last_drop_snap = drop_snap
                    self._fill_table(self.table_drop, drop_snap['rows'])
        except Exception:
            import traceback
            traceback.print_exc()
            loaded = False
        if loaded:
            snap = in_snap or out_snap
            rows = (in_snap or {}).get('rows') or []
            date_text = snap.get('date') or datetime.now().strftime('%Y-%m-%d')
            fetch_time = snap.get('fetch_time') or datetime.now().strftime('%H:%M:%S')
            quote_time = snap.get('quote_time') or '—'
            total = snap.get('total') or len(rows)
            self.point_count_changed.emit(len(rows))
            self.meta_lbl.setText(
                f"📅 {date_text}  {period}主力资金 TOP30 | 行情时间: {quote_time} | 抓取: {fetch_time} [缓存回退] | 总样本: {total}"
            )
            self.status_changed.emit('success', f"上次刷新: {fetch_time} [已冻结]")
            self.bottom_status_changed.emit(
                f"个股资金显示缓存数据 {date_text} {fetch_time} | 周期: {period} | 样本 {total}"
            )
        return loaded

    def _fetch(self):
        if self.worker and self.worker.isRunning():
            self._pending_period = self.period
            return
        self._pending_period = None
        period = self.period
        self.status_changed.emit('loading', f'正在获取个股资金({period})...')
        self.worker = DataWorker(fetcher=lambda: _fetch_stock_top(period=period, top_n=30))
        self.worker.data_ready.connect(self._on_data_ready)
        self.worker.error_occurred.connect(self._on_error)
        self.worker.start()
        if not (self.worker_out and self.worker_out.isRunning()):
            self.worker_out = DataWorker(fetcher=lambda: _fetch_stock_top(period=period, top_n=30, ascending=True))
            self.worker_out.data_ready.connect(self._on_out_data_ready)
            self.worker_out.error_occurred.connect(self._on_error)
            self.worker_out.start()
        if not (self.worker_gain and self.worker_gain.isRunning()):
            self.worker_gain = DataWorker(fetcher=lambda: _fetch_stock_gain_top(top_n=30, ascending=False))
            self.worker_gain.data_ready.connect(self._on_gain_data_ready)
            self.worker_gain.error_occurred.connect(self._on_error)
            self.worker_gain.start()
        if not (self.worker_drop and self.worker_drop.isRunning()):
            self.worker_drop = DataWorker(fetcher=lambda: _fetch_stock_gain_top(top_n=30, ascending=True))
            self.worker_drop.data_ready.connect(self._on_drop_data_ready)
            self.worker_drop.error_occurred.connect(self._on_error)
            self.worker_drop.start()

    def _on_data_ready(self, snap):
        self._last_in_snap = snap
        if self._pending_period and self._pending_period != snap.get('period'):
            QTimer.singleShot(0, self._fetch)
            return
        # fallback=True 表示是 push2 失败 -> Tencent 兜底，资金流字段是缓存陈值
        # 不写入 history 缓存（避免污染日级历史）；latest 缓存仍写以便下次启动可用
        is_fallback = bool(snap.get('fallback'))
        save_panel_cache(self._cache_name('in', snap['period']), snap)
        if not is_fallback:
            _save_stock_history_snapshot('in', snap)
        self._fill_table(self.table_in, snap['rows'])
        self.point_count_changed.emit(len(snap['rows']))
        fallback_tag = ' [腾讯兜底·资金流字段陈]' if is_fallback else ''
        self.meta_lbl.setText(
            f"📅 {snap.get('date', _snapshot_date())}  {snap['period']}主力净流入 TOP30 | 行情时间: {snap['quote_time']} | 抓取: {snap['fetch_time']} | 总样本: {snap['total']}{fallback_tag}"
        )
        self.status_changed.emit('success', f"上次刷新: {snap['fetch_time']}{fallback_tag}")
        self.bottom_status_changed.emit(
            f"个股资金TOP30已更新 {snap['fetch_time']} | 周期: {snap['period']} | 样本 {snap['total']}{fallback_tag}"
        )

    def _on_out_data_ready(self, snap):
        self._last_out_snap = snap
        save_panel_cache(self._cache_name('out', snap['period']), snap)
        if not snap.get('fallback'):
            _save_stock_history_snapshot('out', snap)
        self._fill_table(self.table_out, snap['rows'])

    def _fill_table(self, table, rows):
        table.setRowCount(len(rows))
        last_col = len(_HEADERS) - 1
        for r, row in enumerate(rows):
            values = [
                str(row['rank']),
                row['code'],
                row['name'],
                row.get('sector', ''),
                _format_price(row['price'], row['code']),
                _format_pct(row['pct'], signed=True),
                _format_pct(row['turnover']),
                _format_volume_ratio(row['volume_ratio']),
                _format_money(row['amount']),
                _format_money(row.get('main_net')),
                _format_pct(row.get('main_pct')),
                _format_money(row.get('super_net')),
                _format_money(row.get('big_net')),
                _format_money(row.get('mid_net')),
                _format_money(row.get('small_net')),
            ]
            for c, text in enumerate(values):
                item = QTableWidgetItem(text)
                item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
                if c == 5:
                    item.setForeground(self._brush(row['pct']))
                elif c == 7:
                    item.setForeground(self._volume_ratio_brush(row['volume_ratio']))
                elif c >= 9:
                    key = ['main_net', 'main_pct', 'super_net', 'big_net', 'mid_net', 'small_net'][c - 9]
                    item.setForeground(self._brush(row.get(key)))
                elif c == 2:
                    item.setForeground(QBrush(QColor('#ffffff')))
                elif c == 3:
                    item.setForeground(QBrush(QColor('#e05c5c')))
                table.setItem(r, c, item)
            btn = QPushButton('加自选')
            btn.setFixedSize(62, 24)
            btn.setStyleSheet(
                'QPushButton { background-color: #2a3a5e; color: #ddd; border: 1px solid #4a6a8e; '
                'border-radius: 3px; font-size: 10px; padding: 1px 4px; }'
                'QPushButton:hover { background-color: #e94560; color: #fff; border-color: #e94560; }'
            )
            code, name = row['code'], row['name']
            price = float(row.get('price', 0) or 0)
            btn.clicked.connect(lambda checked, c=code, n=name, p=price: self._quick_add_to_watchlist(c, n, p))
            table.setCellWidget(r, last_col, btn)

    def _on_gain_data_ready(self, snap):
        self._last_gain_snap = snap
        save_panel_cache('stock_gain', snap)
        self._fill_table(self.table_gain, snap['rows'])

    def _on_drop_data_ready(self, snap):
        self._last_drop_snap = snap
        save_panel_cache('stock_drop', snap)
        self._fill_table(self.table_drop, snap['rows'])

    def _on_error(self, msg):
        if self._load_cached_period(self.period):
            self.bottom_status_changed.emit(f'个股资金获取失败，显示冻结数据: {msg}')
            return
        self.status_changed.emit('error', '获取失败')
        self.bottom_status_changed.emit(f'个股资金错误: {msg}')

    def _brush(self, value):
        v = _to_float(value)
        if v > 0:
            return QBrush(QColor(RED))
        if v < 0:
            return QBrush(QColor(GREEN))
        return QBrush(QColor(MUTED))

    def _volume_ratio_brush(self, value):
        v = _to_float(value)
        if v >= 1.2:
            return QBrush(QColor(RED))
        if v <= 0.8:
            return QBrush(QColor(GREEN))
        return QBrush(QColor(MUTED))

    # ---- 自选股 ----

    def _make_watchlist_tab(self):
        widget = QWidget()
        layout = QVBoxLayout(widget)
        layout.setContentsMargins(8, 8, 8, 8)
        layout.setSpacing(8)

        # 顶部输入栏
        input_row = QHBoxLayout()
        input_row.setSpacing(6)
        self._wl_input = QLineEdit()
        self._wl_input.setPlaceholderText('输入代码按回车添加（数量留空按默认市值估算）')
        self._wl_input.setFixedHeight(32)
        self._wl_input.setStyleSheet(
            f"QLineEdit {{ background-color: {CHART_BG}; color: #fff; border: 1px solid #444; "
            f"border-radius: 4px; padding: 4px 10px; font-size: 13px; }}"
        )
        self._wl_input.returnPressed.connect(self._add_watchlist_code)
        input_row.addWidget(self._wl_input, stretch=1)

        add_btn = QPushButton('添加')
        add_btn.setFixedSize(60, 32)
        add_btn.clicked.connect(self._add_watchlist_code)
        input_row.addWidget(add_btn)

        to_holding_btn = QPushButton('→ 加入真实持仓')
        to_holding_btn.setFixedHeight(32)
        to_holding_btn.setToolTip('将选中股票加入真实持仓')
        to_holding_btn.setStyleSheet(
            'QPushButton { background-color: #1a3a5a; color: #7acbe8;'
            ' border: 1px solid #3a6a8a; border-radius: 4px; padding: 0 10px; }'
            'QPushButton:hover { background-color: #2a5a8a; color: #fff; }'
        )
        to_holding_btn.clicked.connect(self._on_watchlist_to_holding)
        input_row.addWidget(to_holding_btn)

        import_wl_btn = QPushButton('📥 批量导入')
        import_wl_btn.setFixedHeight(32)
        import_wl_btn.setToolTip('从通达信/东财 .txt 或 Excel .xlsx/.xlsm 批量导入自选股')
        import_wl_btn.setStyleSheet(
            'QPushButton { background-color: #1a3a2a; color: #7aebb8;'
            ' border: 1px solid #3a8a6a; border-radius: 4px; padding: 0 10px; }'
            'QPushButton:hover { background-color: #2a6a4a; color: #fff; }'
        )
        import_wl_btn.clicked.connect(self._on_import_watchlist)
        input_row.addWidget(import_wl_btn)
        layout.addLayout(input_row)

        # 分组管理栏：当前分组下拉 + 管理按钮
        group_row = QHBoxLayout()
        group_row.setSpacing(6)
        group_hint = QLabel('当前分组：')
        group_hint.setStyleSheet(f'color: {MUTED}; font-size: 12px;')
        group_row.addWidget(group_hint)
        self._wl_group_combo = QComboBox()
        self._wl_group_combo.setFixedHeight(28)
        self._wl_group_combo.setMinimumWidth(140)
        self._wl_group_combo.setStyleSheet(_COMBO_STYLE)
        self._reload_group_combo()
        self._wl_group_combo.currentIndexChanged.connect(self._on_wl_group_changed)
        group_row.addWidget(self._wl_group_combo)
        new_group_btn = QPushButton('新建分组')
        rename_group_btn = QPushButton('重命名分组')
        del_group_btn = QPushButton('删除分组')
        for b, cb in (
            (new_group_btn, self._on_new_group),
            (rename_group_btn, self._on_rename_group),
            (del_group_btn, self._on_delete_group),
        ):
            b.setFixedHeight(28)
            b.setStyleSheet(
                'QPushButton { background-color: #1a2a3a; color: #9ab8cc;'
                ' border: 1px solid #3a5a7a; border-radius: 4px; padding: 0 10px; font-size: 12px; }'
                'QPushButton:hover { background-color: #2a4a6a; color: #fff; }'
            )
            b.clicked.connect(cb)
            group_row.addWidget(b)
        group_row.addStretch(1)
        self._wl_total_label = QLabel('当日总收益: --')
        self._wl_total_label.setStyleSheet(
            'color: #ddd; font-size: 13px; font-weight: bold; padding: 0 8px;')
        group_row.addWidget(self._wl_total_label)
        layout.addLayout(group_row)

        self._wl_pnl_formula_label = QLabel(
            f'{_WATCHLIST_DAILY_PNL_FORMULA}；{_WATCHLIST_TOTAL_PNL_FORMULA}'
        )
        self._wl_pnl_formula_label.setWordWrap(True)
        self._wl_pnl_formula_label.setStyleSheet(
            'color: #9ab8cc; font-size: 12px; padding: 0 8px 4px 8px;')
        layout.addWidget(self._wl_pnl_formula_label)

        # 批量操作栏：全选 + 批量迁移 + 批量删除
        batch_row = QHBoxLayout()
        batch_row.setSpacing(6)
        self._wl_select_all = QCheckBox('全选')
        self._wl_select_all.setStyleSheet(f'color: {MUTED}; font-size: 12px;')
        self._wl_select_all.stateChanged.connect(self._on_wl_select_all)
        batch_row.addWidget(self._wl_select_all)
        batch_migrate_btn = QPushButton('批量迁移到…')
        batch_migrate_btn.clicked.connect(self._on_wl_batch_migrate)
        batch_del_btn = QPushButton('批量删除')
        batch_del_btn.clicked.connect(self._on_wl_batch_delete)
        batch_migrate_btn.setStyleSheet(
            'QPushButton { background-color: #1a2a3a; color: #9ab8cc;'
            ' border: 1px solid #3a5a7a; border-radius: 4px; padding: 0 10px; font-size: 12px; }'
            'QPushButton:hover { background-color: #2a4a6a; color: #fff; }'
        )
        batch_migrate_btn.setFixedHeight(28)
        batch_del_btn.setStyleSheet(
            'QPushButton { background-color: #3a1a1a; color: #e84444;'
            ' border: 1px solid #e84444; border-radius: 4px; padding: 0 10px; font-size: 12px; }'
            'QPushButton:hover { background-color: #e84444; color: #fff; }'
        )
        batch_del_btn.setFixedHeight(28)
        batch_row.addWidget(batch_migrate_btn)
        batch_row.addWidget(batch_del_btn)
        batch_row.addStretch(1)
        layout.addLayout(batch_row)

        # 表格
        self.table_watchlist = QTableWidget(0, len(_WATCHLIST_HEADERS))
        self.table_watchlist.setHorizontalHeaderLabels(_WATCHLIST_HEADERS)
        configure_no_truncation(self.table_watchlist)
        self._apply_watchlist_header_tips()
        self.table_watchlist.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.table_watchlist.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.table_watchlist.setSelectionMode(QAbstractItemView.SelectionMode.SingleSelection)
        self.table_watchlist.verticalHeader().setVisible(False)
        self.table_watchlist.setAlternatingRowColors(True)
        self.table_watchlist.setShowGrid(False)
        self.table_watchlist.setStyleSheet(f'''
            QTableWidget {{
                background-color: {CHART_BG};
                alternate-background-color: #1d2d4a;
                color: #dddddd;
                border: none;
                border-radius: 8px;
                font-size: 12px;
                selection-background-color: #2a3a5e;
            }}
            QHeaderView::section {{
                background-color: #111a30;
                color: #aaaaaa;
                border: none;
                padding: 8px 6px;
                font-weight: bold;
            }}
            QTableWidget::item {{
                padding: 6px 4px;
                border: none;
            }}
        ''')
        h = self.table_watchlist.horizontalHeader()
        h.setSectionResizeMode(QHeaderView.ResizeMode.ResizeToContents)
        h.setSectionResizeMode(0, QHeaderView.ResizeMode.Fixed)
        self.table_watchlist.setColumnWidth(0, 84)
        h.setSectionResizeMode(1, QHeaderView.ResizeMode.Stretch)
        h.setSectionResizeMode(len(_WATCHLIST_HEADERS) - 1, QHeaderView.ResizeMode.Fixed)
        self.table_watchlist.setColumnWidth(len(_WATCHLIST_HEADERS) - 1, 110)
        # 列头点击切换排序（涨跌幅 / 主力净额 / 累计盈亏 / 盈亏比例）
        h.setSectionsClickable(True)
        h.sectionClicked.connect(self._on_wl_header_clicked)
        self.table_watchlist.cellDoubleClicked.connect(self._on_watchlist_double_click)
        self.table_watchlist.setContextMenuPolicy(Qt.ContextMenuPolicy.CustomContextMenu)
        self.table_watchlist.customContextMenuRequested.connect(self._on_wl_context_menu)
        layout.addWidget(self.table_watchlist, stretch=1)

        return widget

    # ---- 自选股列头排序 ----

    _WL_SORTABLE_COLS = {
        6: ('pct', '涨跌幅'),
        7: ('main_net', '主力净额(亿)'),
        11: ('cumulative_pnl', '累计盈亏'),
        12: ('cumulative_pnl_pct', '盈亏比例'),
        13: ('daily_pnl', _WATCHLIST_DAILY_PNL_HEADER),
    }

    def _on_wl_header_clicked(self, col: int):
        """点击列头切换排序：涨跌幅 / 主力净额 / 累计盈亏 / 盈亏比例 可排序。"""
        if col not in self._WL_SORTABLE_COLS:
            return
        if self._wl_sort_col == col:
            self._wl_sort_desc = not self._wl_sort_desc  # 降↔升 切换
        else:
            self._wl_sort_col = col
            self._wl_sort_desc = True
        self._refresh_wl_sort_indicators()
        # 立刻重渲当前数据
        self._rerender_watchlist()

    def _refresh_wl_sort_indicators(self):
        """更新表头文字带 ↓ / ↑ 指示器。"""
        labels = list(_WATCHLIST_HEADERS)
        for c, (_, base) in self._WL_SORTABLE_COLS.items():
            if c == self._wl_sort_col:
                arrow = ' ↓' if self._wl_sort_desc else ' ↑'
                labels[c] = base + arrow
            else:
                labels[c] = base
        self.table_watchlist.setHorizontalHeaderLabels(labels)
        self._apply_watchlist_header_tips()

    def _apply_watchlist_header_tips(self):
        for i in range(self.table_watchlist.columnCount()):
            item = self.table_watchlist.horizontalHeaderItem(i)
            if item is None:
                continue
            base = item.text().replace(' ↓', '').replace(' ↑', '')
            tip = _WATCHLIST_HEADER_TIPS.get(base)
            if tip:
                item.setToolTip(tip)

    def _apply_wl_sort(self, rows: list) -> list:
        """按当前排序状态返回行的副本；未排序时原样返回。"""
        if self._wl_sort_col is None:
            return list(rows)
        key_name, _ = self._WL_SORTABLE_COLS[self._wl_sort_col]

        def _key(r):
            v = r.get(key_name)
            try:
                return float(v) if v is not None else float('-inf')
            except (TypeError, ValueError):
                return float('-inf')

        if self._wl_sort_desc:
            return sorted(rows, key=_key, reverse=True)
        # 升序：None 沉底，其余按值升序
        return sorted(rows, key=lambda r: (_key(r) == float('-inf'), _key(r)))

    def _on_watchlist_double_click(self, row, _col):
        """双击自选股行 → 弹出近10日主力净额柱状图。

        以表格第 0 列实际显示的代码为准，避免排序后行索引与原始 snap 错位。
        """
        code_item = self.table_watchlist.item(row, 0)
        if not code_item:
            return
        code = _clean_code(code_item.text())
        name_item = self.table_watchlist.item(row, 1)
        name = name_item.text() if name_item else code
        if not code:
            return
        records = load_watchlist_fund_history(code)
        dlg = _WatchlistFundDialog(code, name, records, parent=self)
        dlg.exec()

    def _on_import_watchlist(self):
        """从通达信/东财 .txt 或 Excel .xlsx/.xlsm 批量导入自选股。

        支持格式：
        1. 通达信 Excel（列头含"代码"/"ETF代码"，可选"名称"）
        2. 东财 Excel（列头含"代码"/"ETF代码"，可选"名称"）
        3. 通达信自选股导出 .txt（制表符分隔，首列代码/名称/...）
        4. 纯文本代码列表（每行一个代码）
        """
        path, _ = QFileDialog.getOpenFileName(
            self, '导入自选股',
            '', '支持格式 (*.txt *.xlsx *.xlsm);;文本文件 (*.txt);;Excel (*.xlsx *.xlsm)',
        )
        if not path:
            return

        def _normalize(raw_code: str) -> str | None:
            c = raw_code.strip().upper()
            # 去掉常见后缀 .SS .SZ .SH .CSI
            for suf in ('.SS', '.SH', '.SZ', '.CSI'):
                if c.endswith(suf):
                    c = c[:-len(suf)]
                    break
            # 去掉 SH/SZ/BJ 前缀
            if c.startswith(('SH', 'SZ', 'BJ')):
                c = c[2:]
            if not c.isdigit():
                return None
            # 5位代码补前导0 → 6位（如 00981 → 000981）
            if len(c) == 5:
                c = '0' + c
            if len(c) != 6:
                return None
            # 校验 A 股/ETF 代码有效范围，避免把日期/金额等6位数字误识别为代码
            return c if _is_valid_stock_code(c) else None

        codes_found: list[str] = []
        names_found: dict[str, str] = {}  # code -> name

        try:
            if path.lower().endswith(('.xlsx', '.xlsm')):
                codes_found, names_found = self._parse_watchlist_xlsx(path, _normalize)
            elif path.lower().endswith('.txt'):
                codes_found, names_found = self._parse_watchlist_txt(path, _normalize)
            else:
                codes_found, names_found = self._parse_watchlist_txt(path, _normalize)
        except Exception as e:
            QMessageBox.critical(self, '读取失败', str(e))
            return

        if not codes_found:
            QMessageBox.information(self, '提示', '未能从文件中识别出任何股票/ETF代码')
            return

        existing_plain = {_clean_code(r['code']) for r in self._watchlist_codes}
        new_codes: list[str] = []
        seen = set(existing_plain)
        for code in codes_found:
            if code in seen:
                continue
            seen.add(code)
            new_codes.append(code)

        quotes: dict = {}
        if new_codes:
            try:
                quotes = quotes_routed(new_codes)
            except Exception:
                quotes = {}

        import_group = self._wl_active_group or self._wl_default_group
        new_records = self._build_watchlist_import_records(
            codes_found, names_found, existing_plain, quotes, group=import_group,
        )
        added = len(new_records)
        self._watchlist_codes.extend(new_records)

        if added:
            save_watchlist(self._watchlist_codes)
            self._fetch_watchlist()

        total = len(codes_found)
        skipped = total - added
        QMessageBox.information(
            self, '导入完成',
            f'共识别 {total} 只，新增 {added} 只，跳过 {skipped} 只（已存在）。',
        )

    @staticmethod
    def _build_watchlist_import_records(
        codes_found: list[str],
        names_found: dict[str, str],
        existing_plain: set[str],
        quotes: dict,
        group: str = '',
        today: str | None = None,
    ) -> list[dict]:
        today = today or datetime.now().strftime('%Y-%m-%d')
        records: list[dict] = []
        seen = set(existing_plain)
        for code in codes_found:
            if code in seen:
                continue
            seen.add(code)
            quote = quotes.get(code, {}) if isinstance(quotes, dict) else {}
            price = _to_float(quote.get('price')) if quote else 0.0
            cost_price = price if price > 0 else None
            records.append({
                'code': code,
                'name': names_found.get(code, '') or quote.get('name', ''),
                'shares': _default_wl_shares(cost_price),
                'cost_price': cost_price,
                'add_price': cost_price,
                'add_date': today if cost_price else None,
                'cumulative_pnl': 0.0,
                'last_pnl_date': None,
                'theme': '',
                'group': group,
            })
        return records

    @staticmethod
    def _parse_watchlist_xlsx(path, _normalize) -> tuple[list[str], dict[str, str]]:
        """解析通达信/东财 Excel 自选股导出文件。"""
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)

        codes: list[str] = []
        names: dict[str, str] = {}
        seen: set[str] = set()

        try:
            for ws in wb.worksheets:
                code_col: int | None = None
                name_col: int | None = None
                header_row = 0

                # 扫描前 10 行找表头（含"代码"/"ETF代码"列）。
                for ri, row in enumerate(ws.iter_rows(max_row=10, values_only=True), start=1):
                    for ci, cell in enumerate(row):
                        val = str(cell or '').strip()
                        if not val:
                            continue
                        if '代码' in val and code_col is None:
                            code_col = ci
                        elif '名称' in val and name_col is None:
                            name_col = ci
                    if code_col is not None:
                        header_row = ri
                        break

                if code_col is not None:
                    # 表头模式：从"代码"/"ETF代码"列提取。
                    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
                        if code_col >= len(row):
                            continue
                        code = _normalize(str(row[code_col] or ''))
                        if not code:
                            continue
                        name = ''
                        if name_col is not None and name_col < len(row):
                            name = str(row[name_col] or '').strip()
                        if code not in seen:
                            seen.add(code)
                            codes.append(code)
                        if name and name != 'None' and not names.get(code):
                            names[code] = name
                else:
                    # 兜底：扫描当前工作表所有单元格找 6 位代码。
                    for row in ws.iter_rows(values_only=True):
                        for cell in row:
                            code = _normalize(str(cell or ''))
                            if not code or code in seen:
                                continue
                            seen.add(code)
                            codes.append(code)
        finally:
            wb.close()

        return codes, names

    @staticmethod
    def _parse_watchlist_txt(path, _normalize) -> tuple[list[str], dict[str, str]]:
        """解析通达信 .txt 自选股导出文件（制表符分隔或纯代码列表）。"""
        text = open(path, encoding='utf-8', errors='replace').read()
        lines = [ln.strip() for ln in text.splitlines() if ln.strip()]

        codes: list[str] = []
        names: dict[str, str] = {}
        seen: set[str] = set()

        # 检测制表符分隔格式（通达信 .txt 通常是 tab 分隔，首列代码，第二列名称）
        tab_lines = [ln for ln in lines if '\t' in ln]
        if tab_lines and len(tab_lines) >= 1:
            # 尝试找到代码列（检视前几行）
            for ln in tab_lines:
                fields = ln.split('\t')
                if not fields:
                    continue
                code = _normalize(fields[0])
                if code and code not in seen:
                    seen.add(code)
                    codes.append(code)
                    if len(fields) >= 2:
                        name = fields[1].strip()
                        if name and name != code:
                            names[code] = name
        else:
            # 纯代码列表：每行一个（可能是 6 位代码或带前缀/后缀）
            for ln in lines:
                code = _normalize(ln)
                if code and code not in seen:
                    seen.add(code)
                    codes.append(code)

        return codes, names

    def _wl_existing_by_group(self) -> dict[str, set]:
        """{分组名: 该组已有代码集}，供添加弹窗按组查重。"""
        d: dict[str, set] = {}
        for rec in self._watchlist_codes:
            g = rec.get('group') or self._wl_default_group
            d.setdefault(g, set()).add(rec['code'])
        return d

    def _add_watchlist_code(self):
        prefill = self._wl_input.text().strip()
        default_group = self._wl_active_group or self._wl_default_group
        dlg = AddWatchlistDialog(self._wl_existing_by_group(), parent=self,
                                 groups=load_watchlist_groups(), default_group=default_group)
        if prefill:
            dlg._code_input.setText(prefill)
        if dlg.exec() != QDialog.DialogCode.Accepted:
            return
        self._wl_input.clear()
        group = self._commit_group(dlg.result_group)
        self._watchlist_codes.append({
            'code': dlg.result_code,
            'shares': dlg.result_shares,
            'cost_price': dlg.result_cost_price,
            'cumulative_pnl': 0.0,
            'last_pnl_date': None,
            'theme': dlg.result_theme,
            'group': group,
        })
        save_watchlist(self._watchlist_codes)
        self._refresh_one_watchlist(dlg.result_code)

    def _refresh_one_watchlist(self, code: str):
        """加入后只刷这一只的行情并入快照，避免全量刷新过载。后台线程拉取，不阻塞 UI。"""
        self.worker_wl_one = DataWorker(fetcher=lambda: _fetch_watchlist_data([code]))
        self.worker_wl_one.data_ready.connect(lambda snap: self._merge_one_watchlist(code, snap))
        self.worker_wl_one.error_occurred.connect(lambda *_: self._fetch_watchlist())
        self.worker_wl_one.start()

    def _merge_one_watchlist(self, code: str, snap):
        base = self._last_watchlist_snap or {'rows': [], 'date': _snapshot_date()}
        base_rows = [r for r in base.get('rows', []) if r.get('code') != code]
        base_rows.extend(snap.get('rows', []))
        base['rows'] = base_rows
        self._on_watchlist_data_ready(base)

    def _commit_group(self, group: str) -> str:
        """确保分组存在（对话框可能新建），刷新下拉，返回归一化后的分组名。"""
        group = (group or '').strip() or self._wl_default_group
        groups = load_watchlist_groups()
        if group not in groups:
            groups.append(group)
            save_watchlist_groups(groups)
            self._reload_group_combo()
        return group

    def _wl_row_target(self, row: int):
        """从表格行取 (code, name, shares, add_price, cur_price)。

        数量列(2)、加入价列(3)、最新价列(4)；'--'/空 → None。
        """
        code_item = self.table_watchlist.item(row, 0)
        if not code_item:
            return None
        name_item = self.table_watchlist.item(row, 1)

        def _num(col):
            it = self.table_watchlist.item(row, col)
            if not it:
                return None
            try:
                return float(it.text().replace(',', '').strip())
            except (ValueError, TypeError):
                return None

        shares_val = _num(2)
        shares = int(shares_val) if shares_val and shares_val > 0 else None
        add_price = _num(3)
        cur_price = _num(4) or 0.0
        return (_clean_code(code_item.text()),
                name_item.text().strip() if name_item else '未知',
                shares, add_price, cur_price)

    def _on_watchlist_to_holding(self):
        """加入真实持仓：勾选多只则批量（重复跳过）；未勾选则用当前行单只。"""
        targets = []
        for r in range(self.table_watchlist.rowCount()):
            it = self.table_watchlist.item(r, 0)
            if it and it.checkState() == Qt.CheckState.Checked:
                t = self._wl_row_target(r)
                if t:
                    targets.append(t)
        if not targets:
            row = self.table_watchlist.currentRow()
            if row < 0:
                QMessageBox.information(self, '提示', '请先勾选或选中要加入持仓的股票')
                return
            t = self._wl_row_target(row)
            if t:
                targets.append(t)

        if len(targets) == 1:
            code, name, shares, add_price, cur_price = targets[0]
            ref_price = add_price if (add_price and add_price > 0) else cur_price
            dlg = _QuickHoldingDialog(code, name, ref_price, parent=self, ref_shares=shares)
            if dlg.exec() != QDialog.DialogCode.Accepted:
                return
            data = dlg.get_data()
            if not data:
                return
            try:
                add_holding(data)
                QMessageBox.information(self, '已加入',
                    f'{data["name"]}（{data["code"]}）已加入真实持仓\n'
                    f'成本价 {_format_price(data["cost_price"], data.get("code"))}  份数 {data["shares"]:,}')
            except ValueError as e:
                QMessageBox.warning(self, '加入失败', str(e))
            return

        added, skipped = 0, []
        for code, name, shares, add_price, cur_price in targets:
            cost = add_price if (add_price and add_price > 0) else cur_price
            try:
                add_holding({'code': code, 'name': name,
                             'cost_price': round(cost, 3), 'shares': shares or 100,
                             'buy_date': datetime.now().strftime('%Y-%m-%d')})
                added += 1
            except ValueError:
                skipped.append(code)
        self._reset_wl_select_all()
        msg = f'已加入 {added} 只到真实持仓'
        if skipped:
            msg += f'；跳过 {len(skipped)} 只（已在持仓）：{", ".join(skipped[:10])}'
        QMessageBox.information(self, '批量加入完成', msg)

    def _find_wl_record(self, code, group=None):
        """按 (code, 分组) 定位记录；group=None 时取首个同代码记录。"""
        tc = _code_to_tencent(code)
        for rec in self._watchlist_codes:
            if _code_to_tencent(rec.get('code', '')) != tc:
                continue
            if group is None or (rec.get('group') or self._wl_default_group) == group:
                return rec
        return None

    def _remove_watchlist_code(self, code, group=None):
        rec = self._find_wl_record(code, group)
        if rec is None:
            return
        self._watchlist_codes.remove(rec)
        save_watchlist(self._watchlist_codes)
        self._rerender_watchlist()

    def _edit_watchlist_record(self, code, group=None):
        rec = self._find_wl_record(code, group)
        if rec is None:
            return
        dlg = EditWatchlistDialog(rec, parent=self, groups=load_watchlist_groups())
        if dlg.exec() != QDialog.DialogCode.Accepted:
            return
        new_group = self._commit_group(dlg.result_group)
        old_group = rec.get('group') or self._wl_default_group
        if new_group != old_group:
            tc = _code_to_tencent(code)
            clash = any(
                r is not rec
                and _code_to_tencent(r.get('code', '')) == tc
                and (r.get('group') or self._wl_default_group) == new_group
                for r in self._watchlist_codes
            )
            if clash:
                QMessageBox.information(self, '提示', f'分组「{new_group}」已存在该代码')
                return
        rec['shares'] = dlg.result_shares
        rec['cost_price'] = dlg.result_cost_price
        rec['theme'] = dlg.result_theme
        rec['group'] = new_group
        save_watchlist(self._watchlist_codes)
        self._fetch_watchlist()

    def _on_flow_context_menu(self, pos, table):
        """资金流表格右键菜单：加自选。"""
        row_idx = table.rowAt(pos.y())
        if row_idx < 0:
            return
        code_item = table.item(row_idx, 1)
        name_item = table.item(row_idx, 2)
        price_item = table.item(row_idx, 4)
        if not code_item:
            return
        code = code_item.text().strip()
        name = name_item.text().strip() if name_item else ''
        try:
            price = float(price_item.text() or 0) if price_item else 0.0
        except (ValueError, TypeError):
            price = 0.0
        menu = QMenu(self)
        menu.setStyleSheet(
            'QMenu { background-color: #1e2a3a; color: #ddd; border: 1px solid #4a6a8e; padding: 4px; }'
            'QMenu::item { padding: 4px 20px; }'
            'QMenu::item:selected { background-color: #e94560; }'
        )
        add_action = menu.addAction(f'加自选 — {code} {name}')
        action = menu.exec(table.viewport().mapToGlobal(pos))
        if action == add_action:
            self._quick_add_to_watchlist(code, name, price)

    def _quick_add_to_watchlist(self, code: str, name: str, price: float = 0.0):
        """快速加自选：从资金流表或板块强势股调用。"""
        default_group = self._wl_active_group or self._wl_default_group
        tc = _code_to_tencent(code)
        existing = next((r for r in self._watchlist_codes
                         if _code_to_tencent(r['code']) == tc
                         and (r.get('group') or self._wl_default_group) == default_group), None)
        if existing:
            reply = QMessageBox.question(
                self, '已在自选',
                f'{_clean_code(code)} {name} 已在自选列表中。\n是否打开编辑/追加买卖？',
                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
            )
            if reply == QMessageBox.StandardButton.Yes:
                cur_price = price
                if cur_price <= 0:
                    try:
                        quotes = quotes_routed([code])
                        q = quotes.get(code, {})
                        cur_price = float(q.get('price', 0) or 0)
                    except Exception:
                        pass
                dlg = EditTradeDialog(existing, parent=self, current_price=cur_price)
                if dlg.exec() == QDialog.DialogCode.Accepted:
                    r = dlg.get_result()
                    if r:
                        existing['shares'] = r['new_shares']
                        existing['cost_price'] = r['new_cost']
                        existing['theme'] = r.get('theme') or existing.get('theme', '')
                        if r.get('trade_record'):
                            existing.setdefault('trades', []).append(r['trade_record'])
                        save_watchlist(self._watchlist_codes)
                        self._fetch_watchlist()
            return

        dlg = AddWatchlistDialog(self._wl_existing_by_group(), parent=self,
                                 prefill_code=code, prefill_name=name, prefill_price=price,
                                 groups=load_watchlist_groups(), default_group=default_group)
        if dlg.exec() != QDialog.DialogCode.Accepted:
            return
        bd = dlg.result_buy_date or datetime.now().strftime('%Y-%m-%d')
        group = self._commit_group(dlg.result_group)
        rec = {
            'code': dlg.result_code,
            'name': dlg.result_name or name,
            'shares': dlg.result_shares,
            'cost_price': dlg.result_cost_price,
            'cumulative_pnl': 0.0,
            'last_pnl_date': None,
            'theme': dlg.result_theme,
            'group': group,
            'buy_date': bd,
            'trades': [{
                'date': bd, 'direction': 'buy',
                'shares': dlg.result_shares,
                'price': dlg.result_cost_price,
                'timestamp': datetime.now().isoformat(),
                'source': 'quick_add',
            }],
            'pre_close': None, 'last_close': None, 'last_close_date': None,
            'include_in_pnl': True,
        }
        self._watchlist_codes.append(rec)
        save_watchlist(self._watchlist_codes)
        self._fetch_watchlist()

    def _fetch_watchlist(self):
        if not self._watchlist_codes:
            self._last_watchlist_snap = None
            self._fill_watchlist_table([])
            return
        if self.worker_watchlist and self.worker_watchlist.isRunning():
            return
        codes = list(dict.fromkeys(r['code'] for r in self._watchlist_codes))
        self.worker_watchlist = DataWorker(fetcher=lambda: _fetch_watchlist_data(codes))
        self.worker_watchlist.data_ready.connect(self._on_watchlist_data_ready)
        self.worker_watchlist.error_occurred.connect(self._on_watchlist_error)
        self.worker_watchlist.start()

    def post_close_refresh(self):
        """盘后冻结：错开依次刷新自选与期权，避免与持仓监控同时拉取造成过载。"""
        QTimer.singleShot(2000, self._fetch_watchlist)
        QTimer.singleShot(8000, lambda: self._option_refresh_quotes(silent=True, force=True))

    def _build_wl_display_rows(self, snap) -> list[dict]:
        """按自选记录生成显示行（每条记录一行，行带 group）。

        同一代码可在多个分组各有一条记录，分别拥有独立的股数/加入价/盈亏；
        行情按代码共享（从快照取并复制，避免互相覆盖）。
        """
        quote_map = {}
        for q in (snap or {}).get('rows', []):
            quote_map[_code_to_tencent(q['code'])] = q

        today = datetime.now().strftime('%Y-%m-%d')
        wl_updated = False
        rows: list[dict] = []
        for rec in self._watchlist_codes:
            code = rec.get('code')
            if not code:
                continue
            q = quote_map.get(_code_to_tencent(code))
            row = dict(q) if q else {
                'code': code, 'name': rec.get('name', ''),
                'price': None, 'pre_close': None, 'pct': None,
            }

            shares = rec.get('shares')
            cur_price = float(row.get('price') or 0)
            daily_shares = shares if shares else _default_wl_shares(cur_price)

            # 加入价 + 加入日：首次拿到现价时锁定（用户填了加入价则优先用用户的）
            if cur_price > 0:
                if not rec.get('add_price'):
                    rec['add_price'] = rec.get('cost_price') or cur_price
                    wl_updated = True
                if not rec.get('add_date'):
                    rec['add_date'] = today
                    wl_updated = True

            base_price = rec.get('cost_price') or rec.get('add_price')
            add_date = rec.get('add_date')
            row['shares'] = shares
            row['daily_shares'] = daily_shares
            row['cost_price'] = base_price
            row['add_price'] = base_price
            # 累计盈亏 = (现价 − 加入价) × 数量；加入当天锁空，次日起生效
            if base_price and daily_shares and cur_price > 0 and add_date and add_date != today:
                row['cumulative_pnl'] = round((cur_price - base_price) * daily_shares, 2)
                row['cumulative_pnl_pct'] = (cur_price - base_price) / base_price * 100
            else:
                row['cumulative_pnl'] = None
                row['cumulative_pnl_pct'] = None
            # 当日收益 = (现价 − 昨收) × 数量
            pre_close = float(row.get('pre_close') or 0)
            if daily_shares and cur_price > 0 and pre_close > 0:
                row['daily_pnl'] = round((cur_price - pre_close) * daily_shares, 2)
            else:
                row['daily_pnl'] = None

            row['group'] = rec.get('group') or self._wl_default_group
            rows.append(row)

        if wl_updated:
            save_watchlist(self._watchlist_codes)
        return rows

    def _on_watchlist_data_ready(self, snap):
        if not self._watchlist_codes:
            self._last_watchlist_snap = None
            self._fill_watchlist_table([])
            return
        # 如果 push2 部分失败，用缓存补充资金字段
        cached = self._last_watchlist_snap
        if cached and cached.get('rows'):
            cached_map = {r['code']: r for r in cached['rows']}
            for row in snap['rows']:
                if not row.get('has_fund') and row['code'] in cached_map:
                    cr = cached_map[row['code']]
                    row['main_net'] = cr.get('main_net')
                    row['main_pct'] = cr.get('main_pct')
                    row['super_net'] = cr.get('super_net')
                    row['big_net'] = cr.get('big_net')

        self._last_watchlist_snap = snap
        save_panel_cache('watchlist_data', snap)
        self._fill_watchlist_table(self._build_wl_display_rows(snap))

        # 自选股涨跌幅提醒
        if is_trading_time():
            self._check_watchlist_alerts(snap['rows'])

        # 写入自选股逐日资金历史（仅 push2 成功的行，避免污染）
        date_key = snap.get('date') or _snapshot_date()
        for row in snap.get('rows', []):
            if not row.get('has_fund'):
                continue
            code = _clean_code(row['code'])
            save_watchlist_fund_history(code, {
                'date': date_key,
                'name': row.get('name', ''),
                'main_net': row.get('main_net'),
                'main_pct': row.get('main_pct'),
                'super_net': row.get('super_net'),
                'big_net': row.get('big_net'),
            })

    _OPTION_HEADERS = [
        '代码', '简称', '标的', '类型', '持仓方向', '行权价', '到期日', '剩余天数',
        '合约单位', '张数', '名义敞口', '开仓价', '当前价', '前结市值', '浮盈亏', '状态', '操作',
    ]

    def _setup_option_tab(self):
        w = QWidget()
        layout = QVBoxLayout(w)
        layout.setContentsMargins(8, 4, 8, 4)
        layout.setSpacing(6)

        hint = QLabel('可手动维护，也可尝试 AKShare 免费行情刷新。请维护合约代码、标的、到期日、开仓价和当前价。')
        hint.setAlignment(Qt.AlignmentFlag.AlignCenter)
        hint.setStyleSheet('color: #f0c040; font-size: 12px; padding: 4px;')
        layout.addWidget(hint)

        self._option_summary_label = QLabel()
        self._option_summary_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self._option_summary_label.setStyleSheet(
            'color: #ddd; font-size: 12px; padding: 4px 8px; '
            'background-color: #1d2d4a; border-radius: 4px;'
        )
        layout.addWidget(self._option_summary_label)

        input_row = QHBoxLayout()
        input_row.setSpacing(6)
        self._opt_code_input = QLineEdit()
        self._opt_code_input.setPlaceholderText('输入8位代码或完整合约简称…')
        self._opt_code_input.setStyleSheet(_LINE_EDIT_STYLE)
        input_row.addWidget(self._opt_code_input, stretch=1)
        add_btn = QPushButton('添加')
        add_btn.setFixedWidth(60)
        add_btn.setStyleSheet(
            'QPushButton { background-color: #2a6a3a; color: #fff; border: 1px solid #3a8a4a; '
            'border-radius: 4px; padding: 4px 12px; font-size: 12px; }'
            'QPushButton:hover { background-color: #3a8a4a; }'
        )
        add_btn.clicked.connect(self._option_add)
        input_row.addWidget(add_btn)

        refresh_btn = QPushButton('刷新行情')
        refresh_btn.setFixedWidth(72)
        refresh_btn.setStyleSheet(
            'QPushButton { background-color: #2a4a6a; color: #fff; border: 1px solid #3a6a9a; '
            'border-radius: 4px; padding: 4px 8px; font-size: 12px; }'
            'QPushButton:hover { background-color: #3a6a9a; }'
        )
        refresh_btn.clicked.connect(self._option_refresh_quotes)
        input_row.addWidget(refresh_btn)
        input_row.addStretch(1)
        self._option_total_label = QLabel('当日总盈亏: --')
        self._option_total_label.setStyleSheet(
            'color: #ddd; font-size: 14px; font-weight: bold; padding: 0 8px;')
        input_row.addWidget(self._option_total_label)
        layout.addLayout(input_row)

        n_cols = len(self._OPTION_HEADERS)
        self._option_table = QTableWidget(0, n_cols)
        self._option_table.setHorizontalHeaderLabels(self._OPTION_HEADERS)
        configure_no_truncation(self._option_table)
        self._option_table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self._option_table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self._option_table.setSelectionMode(QAbstractItemView.SelectionMode.SingleSelection)
        self._option_table.verticalHeader().setVisible(False)
        self._option_table.setAlternatingRowColors(True)
        self._option_table.setShowGrid(False)
        self._option_table.setStyleSheet(f'''
            QTableWidget {{
                background-color: {CHART_BG};
                alternate-background-color: #1d2d4a;
                color: #dddddd;
                border: none;
                border-radius: 8px;
                font-size: 11px;
            }}
            QHeaderView::section {{
                background-color: #111a30;
                color: #aaaaaa;
                border: none;
                padding: 8px 4px;
                font-weight: bold;
            }}
            QTableWidget::item {{
                padding: 5px 4px;
                border: none;
            }}
        ''')
        oh = self._option_table.horizontalHeader()
        oh.setSectionResizeMode(QHeaderView.ResizeMode.ResizeToContents)
        oh.setSectionResizeMode(1, QHeaderView.ResizeMode.Stretch)
        last_opt_col = n_cols - 1
        oh.setSectionResizeMode(last_opt_col, QHeaderView.ResizeMode.Fixed)
        self._option_table.setColumnWidth(last_opt_col, 130)
        layout.addWidget(self._option_table, stretch=1)

        self._refresh_option_table()
        return w

    def _refresh_option_table(self):
        records = load_option_watchlist()
        t = self._option_table
        t.setRowCount(len(records))
        last_col = len(self._OPTION_HEADERS) - 1

        holding_count = 0
        total_premium_cost = 0.0
        total_market_value = 0.0
        total_floating_pnl = 0.0
        total_notional_exposure = 0.0
        near_expiry_7d = 0
        already_expired = 0

        opt_type_map = {'call': '认购', 'put': '认沽'}
        pos_side_map = {'long_right': '买入权利仓', 'short_obligation': '卖出义务仓'}
        status_map = {
            'holding': '持有中', 'closed': '已平仓',
            'expired': '已到期', 'exercised': '已行权', 'void': '已作废',
        }

        for r, rec in enumerate(records):
            metrics = compute_option_metrics(rec)
            code = rec.get('code', '')
            name = rec.get('name', '')
            underlying_code = rec.get('underlying_code', '')
            underlying_name = rec.get('underlying_name', '')
            option_type = rec.get('option_type', 'call')
            position_side = rec.get('position_side', 'long_right')
            strike_price = rec.get('strike_price', 0.0)
            expiry_date = rec.get('expiry_date', '')
            contract_unit = rec.get('contract_unit', 10000)
            contracts = rec.get('contracts', 0)
            open_price = rec.get('open_price', 0.0)
            current_price = rec.get('current_price', 0.0)
            status = rec.get('status', 'holding')

            if underlying_name and underlying_code:
                underlying_display = f'{underlying_name}({underlying_code})'
            elif underlying_name:
                underlying_display = underlying_name
            elif underlying_code:
                underlying_display = underlying_code
            else:
                underlying_display = '—'

            opt_type_display = opt_type_map.get(option_type, option_type)
            pos_side_display = pos_side_map.get(position_side, position_side)
            status_display = status_map.get(status, status)

            days = metrics['days_to_expiry']
            expiry_level = metrics['expiry_level']
            if status == 'holding' and days is not None and days < 0:
                days_display = '已到期'
            elif days is None:
                days_display = '—'
            else:
                days_display = str(days)

            if metrics['notional_exposure'] is not None:
                exposure_display = f'{metrics["notional_exposure"]:,.0f}'
            else:
                exposure_display = '—'

            if metrics['premium_cost'] is not None:
                cost_display = f'{metrics["premium_cost"]:,.3f}'
            elif metrics['unsupported_short']:
                cost_display = '暂不计算'
            else:
                cost_display = '—'

            pnl_color = None
            if metrics['floating_pnl'] is not None:
                pnl_val = metrics['floating_pnl']
                pnl_display = f'{pnl_val:+,.3f}'
                pnl_color = '#e84444' if pnl_val >= 0 else '#2ecc71'
            elif metrics['needs_price']:
                pnl_display = '需维护'
            elif metrics['unsupported_short']:
                pnl_display = '暂不计算'
            else:
                pnl_display = '—'

            values = [
                code, name or '—', underlying_display,
                opt_type_display, pos_side_display,
                _format_price(strike_price, is_option=True) if strike_price else '—',
                expiry_date or '—', days_display,
                str(contract_unit), str(contracts), exposure_display,
                _format_price(open_price, is_option=True) if open_price else '—',
                _format_price(current_price, is_option=True) if current_price > 0 else ('未维护' if current_price == 0 else '—'),
                cost_display, pnl_display, status_display,
            ]

            for c, val in enumerate(values):
                item = QTableWidgetItem(str(val))
                item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
                t.setItem(r, c, item)

            days_item = t.item(r, 7)
            if days_item:
                if expiry_level == 'expired' or days_display == '已到期':
                    days_item.setForeground(QColor('#777777'))
                elif expiry_level == 'd1':
                    days_item.setForeground(QColor('#e84444'))
                elif expiry_level == 'd3':
                    days_item.setForeground(QColor('#ff8c00'))
                elif expiry_level == 'd7':
                    days_item.setForeground(QColor('#f0c040'))

            if pnl_color:
                pnl_item = t.item(r, 14)
                if pnl_item:
                    pnl_item.setForeground(QColor(pnl_color))
            elif pnl_display == '需维护':
                pnl_item = t.item(r, 14)
                if pnl_item:
                    pnl_item.setForeground(QColor('#f0c040'))
            elif pnl_display == '暂不计算':
                pnl_item = t.item(r, 14)
                if pnl_item:
                    pnl_item.setForeground(QColor('#aaaaaa'))

            if status == 'holding':
                holding_count += 1
                if metrics['notional_exposure'] is not None:
                    total_notional_exposure += metrics['notional_exposure']
                if metrics['premium_cost'] is not None:
                    total_premium_cost += metrics['premium_cost']
                if metrics['market_value'] is not None:
                    total_market_value += metrics['market_value']
                if days is not None and days >= 0 and metrics['floating_pnl'] is not None:
                    total_floating_pnl += metrics['floating_pnl']
                if days is not None:
                    if days < 0:
                        already_expired += 1
                    elif days <= 7:
                        near_expiry_7d += 1

            ops_widget = QWidget()
            ops_widget.setStyleSheet('background: transparent;')
            ops_layout = QHBoxLayout(ops_widget)
            ops_layout.setContentsMargins(2, 2, 2, 2)
            ops_layout.setSpacing(4)

            edit_btn = QPushButton('编辑')
            edit_btn.setFixedSize(54, 26)
            edit_btn.setStyleSheet(
                'QPushButton { background-color: #2a3a5e; color: #ddd; border: 1px solid #4a6a8e; '
                'border-radius: 3px; font-size: 11px; padding: 2px 4px; }'
                'QPushButton:hover { background-color: #e94560; color: #fff; }'
            )
            edit_btn.clicked.connect(lambda checked, ri=r: self._option_edit(ri))
            ops_layout.addWidget(edit_btn)

            del_btn = QPushButton('删除')
            del_btn.setFixedSize(54, 26)
            del_btn.setStyleSheet(
                'QPushButton { background-color: #3a2a2a; color: #ddd; border: 1px solid #6a3a3a; '
                'border-radius: 3px; font-size: 11px; padding: 2px 4px; }'
                'QPushButton:hover { background-color: #c0392b; color: #fff; }'
            )
            del_btn.clicked.connect(lambda checked, ri=r: self._option_delete(ri))
            ops_layout.addWidget(del_btn)

            t.setCellWidget(r, last_col, ops_widget)

        pnl_sign = '+' if total_floating_pnl >= 0 else ''
        _ot_color = '#39c87a' if total_floating_pnl > 0 else ('#e05c5c' if total_floating_pnl < 0 else '#ddd')
        self._option_total_label.setText(
            f'当日总盈亏: <span style="color:{_ot_color}">{pnl_sign}{total_floating_pnl:,.2f}</span>')
        self._option_summary_label.setText(
            f'持有合约数: {holding_count}  |  '
            f'总名义敞口: {total_notional_exposure:,.0f}  |  '
            f'前结市值: {total_premium_cost:,.3f}  |  '
            f'当前总市值: {total_market_value:,.3f}  |  '
            f'合计当日收益: {pnl_sign}{total_floating_pnl:,.2f}  |  '
            f'7天内到期: {near_expiry_7d}  |  '
            f'已到期未处理: {already_expired}'
        )

    def _option_add(self):
        code = self._opt_code_input.text().strip()
        dlg = _OptionPositionDialog(self)
        if code:
            dlg._code_edit.setText(code)
        if dlg.exec() != QDialog.DialogCode.Accepted:
            return
        new_rec = dlg.get_record()
        records = load_option_watchlist()
        records.append(new_rec)
        save_option_watchlist(records)
        self._opt_code_input.clear()
        self._refresh_option_views()

    def _option_edit(self, row_idx):
        records = load_option_watchlist()
        if row_idx < 0 or row_idx >= len(records):
            return
        rec = records[row_idx]
        dlg = _OptionPositionDialog(self, rec)
        if dlg.exec() != QDialog.DialogCode.Accepted:
            return
        updated = dlg.get_record()
        records[row_idx] = updated
        save_option_watchlist(records)
        self._refresh_option_views()

    def _option_delete(self, row_idx):
        records = load_option_watchlist()
        if row_idx < 0 or row_idx >= len(records):
            return
        rec = records[row_idx]
        reply = QMessageBox.question(
            self, '确认删除', f'确定删除期权 {rec.get("code", "")} 吗？',
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
        )
        if reply != QMessageBox.StandardButton.Yes:
            return
        records.pop(row_idx)
        save_option_watchlist(records)
        self._refresh_option_views()

    def _refresh_option_views(self):
        self._refresh_option_table()
        if self._last_watchlist_snap:
            self._rerender_watchlist()

    def _option_refresh_quotes(self, silent=False, force=False):
        """通过 AKShare 免费接口刷新期权持仓的当前价和元数据。

        Args:
            silent: True=自动/切 Tab 触发，不弹窗；force=False 时仍按交易时段/5 分钟节流
            force: True=期权 8 秒自动刷新、切页刷新、盘后冻结等场景，绕过交易时段/节流门控
        """
        if silent and not force:
            if not is_trading_time():
                return
            if time.time() - self._last_option_refresh < 300:
                return

        records = load_option_watchlist()
        holding_records = [(i, r) for i, r in enumerate(records) if r.get('status', 'holding') == 'holding']
        if not holding_records:
            if not silent:
                QMessageBox.information(self, '提示', '没有需要刷新行情的持有中期权合约。')
            return

        codes = [r.get('code', '') for _, r in holding_records]
        codes = [c for c in codes if c and len(c) == 8]
        if not codes:
            if not silent:
                QMessageBox.information(self, '提示', '没有有效的 8 位期权合约代码。')
            return

        refresh_btn = self.sender() if not silent else None
        if getattr(self, '_option_worker', None) and self._option_worker.isRunning():
            if not silent:
                QMessageBox.information(self, '正在刷新', '期权行情正在刷新中，请稍后再试。')
            return

        self._last_option_refresh = time.time()

        if refresh_btn:
            refresh_btn.setEnabled(False)
            refresh_btn.setText('刷新中…')

        self._option_worker = DataWorker(fetcher=lambda: fetch_option_quotes(codes, timeout=10.0))
        self._option_worker.data_ready.connect(
            lambda results: self._on_option_quotes_ready(
                results, holding_records, records, silent, refresh_btn))
        self._option_worker.error_occurred.connect(
            lambda msg: self._on_option_quotes_error(msg, silent, refresh_btn))
        self._option_worker.start()

    def _on_option_quotes_error(self, msg, silent, refresh_btn):
        if refresh_btn:
            refresh_btn.setEnabled(True)
            refresh_btn.setText('刷新行情')
        if not silent:
            QMessageBox.warning(self, '刷新失败', f'期权行情获取异常：{msg}')

    def _on_option_quotes_ready(self, results, holding_records, records, silent, refresh_btn):
        current_records = load_option_watchlist()
        if not current_records:
            if refresh_btn:
                refresh_btn.setEnabled(True)
                refresh_btn.setText('刷新行情')
            self._refresh_option_table()
            return
        updated = 0
        failed_codes = []
        stale_codes = []
        requested_codes = {
            str(rec.get('code', '')).strip()
            for _, rec in holding_records
            if rec.get('code')
        }
        for idx, rec in enumerate(current_records):
            if rec.get('status', 'holding') != 'holding':
                continue
            code = rec.get('code', '')
            if requested_codes and code not in requested_codes:
                continue
            q = results.get(code)
            if q is None or not q.get('ok'):
                if code:
                    failed_codes.append(code)
                continue
            changed = False

            # 真·已到期（到期日早于今天）才保留手填价；盘前/未开盘只是非当日，照常覆盖
            expiry_ref = q.get('expiry_date') or rec.get('expiry_date') or ''
            truly_expired = bool(expiry_ref) and expiry_ref[:10] < datetime.now().strftime('%Y-%m-%d')

            if truly_expired:
                if code:
                    stale_codes.append(code)
            else:
                price = q.get('current_price', 0.0)
                if price and price > 0:
                    rec['current_price'] = price
                    changed = True
                pre = q.get('pre_close', 0.0)
                if pre and pre > 0:
                    rec['pre_close'] = pre
                    changed = True

            # 合约规格以交易所实时数据为准覆盖（持仓方向除外，纯手填）
            if q.get('name') and rec.get('name') != q['name']:
                rec['name'] = q['name']
                changed = True
            if q.get('underlying_code') and rec.get('underlying_code') != q['underlying_code']:
                rec['underlying_code'] = q['underlying_code']
                changed = True
            if q.get('strike_price', 0) > 0 and rec.get('strike_price') != q['strike_price']:
                rec['strike_price'] = q['strike_price']
                changed = True
            if q.get('option_type') and rec.get('option_type') != q['option_type']:
                rec['option_type'] = q['option_type']
                changed = True
            if q.get('contract_unit', 0) > 0 and rec.get('contract_unit') != q['contract_unit']:
                rec['contract_unit'] = q['contract_unit']
                changed = True
            if q.get('expiry_date') and rec.get('expiry_date') != q['expiry_date']:
                rec['expiry_date'] = q['expiry_date']
                changed = True

            if changed:
                rec['updated_at'] = datetime.now().isoformat()
                current_records[idx] = rec
                updated += 1

        if updated:
            save_option_watchlist(current_records)
        self._refresh_option_table()
        if self._last_watchlist_snap:
            self._rerender_watchlist()

        if refresh_btn:
            refresh_btn.setEnabled(True)
            refresh_btn.setText('刷新行情')

        if silent:
            return

        parts = [f'成功更新 {updated} 条期权行情。']
        if stale_codes:
            parts.append(
                f'以下 {len(stale_codes)} 条为已到期合约（保留手填价，未覆盖）：'
                f'{", ".join(stale_codes[:10])}'
            )
        if failed_codes:
            parts.append(
                f'以下 {len(failed_codes)} 条未获取到行情（原手填价保留）：'
                f'{", ".join(failed_codes[:10])}'
            )
        if stale_codes or failed_codes:
            QMessageBox.warning(self, '刷新完成（部分未更新）', '\n'.join(parts))
        else:
            QMessageBox.information(self, '刷新完成', parts[0])

    def _check_watchlist_alerts(self, rows):
        """检查自选股涨跌幅是否触及阈值，触发分级提醒。"""
        ac = getattr(self, 'alert_center', None)
        if not ac or not ac.get_setting('watchlist_pct_enabled', True):
            return
        thresholds = sorted(ac.get_setting('watchlist_pct_thresholds', [5, 8, 10]), reverse=True)
        cooldown = ac.get_setting('watchlist_pct_cooldown', 600)
        for row in rows:
            code = _clean_code(row.get('code', ''))
            name = row.get('name', '') or code
            try:
                pct = float(row.get('pct', 0) or 0)
            except (TypeError, ValueError):
                continue
            if pct == 0:
                continue
            direction = 'up' if pct > 0 else 'dn'
            abs_pct = abs(pct)
            for th in thresholds:
                if abs_pct >= th:
                    key = f'watchlist_pct:{code}:{direction}:{th}'
                    emoji = '🚀' if pct > 0 else '📉'
                    level = 'critical' if th >= 10 else 'warning' if th >= 8 else 'info'
                    ac.try_trigger(
                        event_type='watchlist_pct',
                        title=f'{emoji} 自选股: {name}',
                        message=f'{name}({code}) {pct:+.2f}%，触及{"+" if pct > 0 else ""}{th}%档',
                        level=level,
                        key=key,
                        cooldown=cooldown,
                    )
                    break  # 只取最高档

    def _on_watchlist_error(self, msg):
        if self._load_cached_watchlist():
            return
        self.bottom_status_changed.emit(f'自选股获取失败: {msg}')

    def _load_cached_watchlist(self):
        try:
            if not self._watchlist_codes:
                self._last_watchlist_snap = None
                return False
            snap = load_panel_cache('watchlist_data')
            if snap and snap.get('rows'):
                self._last_watchlist_snap = snap
                self._fill_watchlist_table(self._build_wl_display_rows(snap))
                return True
        except Exception:
            import traceback
            traceback.print_exc()
        return False

    def _ensure_watchlist_groups(self):
        """保证：至少有一个分组，且每只自选股都归属于一个有效分组。

        首次启用分组时，现有股票（group 为空）统一归入「新建分组1」。
        """
        groups = load_watchlist_groups()
        changed_g = False
        if not groups:
            groups = [self._wl_default_group]
            changed_g = True
        valid = set(groups)
        changed_w = False
        for rec in self._watchlist_codes:
            g = rec.get('group') or ''
            if g not in valid:
                rec['group'] = groups[0]
                changed_w = True
        if changed_g:
            save_watchlist_groups(groups)
        if changed_w:
            save_watchlist(self._watchlist_codes)

    def _reload_group_combo(self):
        """重建分组下拉：全部 + 各分组。尽量保持当前选中项。"""
        combo = self._wl_group_combo
        prev = combo.currentText() if combo.count() else '全部'
        combo.blockSignals(True)
        combo.clear()
        combo.addItem('全部')
        for g in load_watchlist_groups():
            combo.addItem(g)
        idx = combo.findText(prev)
        combo.setCurrentIndex(idx if idx >= 0 else 0)
        combo.blockSignals(False)

    def _on_wl_group_changed(self, _idx: int):
        text = self._wl_group_combo.currentText()
        self._wl_active_group = None if text == '全部' else text
        self._rerender_watchlist()

    def _fill_watchlist_table(self, rows):
        rows = self._apply_wl_sort(rows)
        # 按当前选中分组过滤（全部时不过滤）；行自带 group
        if self._wl_active_group is not None:
            rows = [r for r in rows
                    if (r.get('group') or self._wl_default_group) == self._wl_active_group]

        checked_keys = self._wl_checked_keys()
        self.table_watchlist.clearSpans()
        self.table_watchlist.setRowCount(len(rows))
        for r, row in enumerate(rows):
            self._render_wl_row(r, row, checked_keys)
        self._update_wl_total_label(rows)

    def _update_wl_total_label(self, rows):
        """当前分组当日总收益 = Σ本组个股当日收益 + 期权合计当日收益。"""
        stock_total = sum(r.get('daily_pnl') or 0 for r in rows)
        try:
            from core.cache import compute_option_total_daily_pnl
            opt_total = compute_option_total_daily_pnl()
        except Exception:
            opt_total = 0.0
        total = stock_total + opt_total
        sign = '+' if total >= 0 else ''
        color = '#39c87a' if total > 0 else ('#e05c5c' if total < 0 else '#ddd')
        self._wl_total_label.setText(
            f'当日总收益: <span style="color:{color}">{sign}{total:,.2f}</span>'
            f'　(个股 {stock_total:+,.2f} / 期权 {opt_total:+,.2f})')

    def _rerender_watchlist(self):
        """用上一次快照重渲当前分组视图。"""
        self._fill_watchlist_table(
            self._build_wl_display_rows(self._last_watchlist_snap))

    # ---- 分组管理 ----

    def _on_new_group(self):
        name, ok = QInputDialog.getText(self, '新建分组', '分组名称：')
        name = (name or '').strip()
        if not ok or not name:
            return
        groups = load_watchlist_groups()
        if name in groups:
            QMessageBox.information(self, '提示', f'分组「{name}」已存在')
            return
        groups.append(name)
        save_watchlist_groups(groups)
        self._reload_group_combo()
        # 切到新建分组
        idx = self._wl_group_combo.findText(name)
        if idx >= 0:
            self._wl_group_combo.setCurrentIndex(idx)

    def _on_rename_group(self):
        groups = load_watchlist_groups()
        if not groups:
            QMessageBox.information(self, '提示', '暂无可重命名的分组')
            return
        old, ok = QInputDialog.getItem(self, '重命名分组', '选择分组：', groups, 0, False)
        if not ok or not old:
            return
        new, ok = QInputDialog.getText(self, '重命名分组', f'将「{old}」改名为：')
        new = (new or '').strip()
        if not ok or not new or new == old:
            return
        if new in groups:
            QMessageBox.information(self, '提示', f'分组「{new}」已存在')
            return
        groups[groups.index(old)] = new
        save_watchlist_groups(groups)
        for rec in self._watchlist_codes:
            if (rec.get('group') or '') == old:
                rec['group'] = new
        save_watchlist(self._watchlist_codes)
        if self._wl_active_group == old:
            self._wl_active_group = new
        self._reload_group_combo()
        self._rerender_watchlist()

    def _on_delete_group(self):
        groups = load_watchlist_groups()
        if len(groups) <= 1:
            QMessageBox.information(self, '提示', '至少保留一个分组，无法删除')
            return
        name, ok = QInputDialog.getItem(
            self, '删除分组', '选择分组（其中股票将移到第一个分组）：', groups, 0, False)
        if not ok or not name:
            return
        groups.remove(name)
        save_watchlist_groups(groups)
        target = groups[0]
        for rec in self._watchlist_codes:
            if (rec.get('group') or '') == name:
                rec['group'] = target
        save_watchlist(self._watchlist_codes)
        if self._wl_active_group == name:
            self._wl_active_group = None
        self._reload_group_combo()
        self._rerender_watchlist()

    def _on_wl_context_menu(self, pos):
        row = self.table_watchlist.rowAt(pos.y())
        if row < 0:
            return
        code_item = self.table_watchlist.item(row, 0)
        if not code_item:
            return
        code = _clean_code(code_item.text())
        if not code:
            return
        src_group = code_item.data(Qt.ItemDataRole.UserRole)
        menu = QMenu(self)
        menu.setStyleSheet(
            'QMenu { background-color: #1e2a3a; color: #ddd; border: 1px solid #4a6a8e; padding: 4px; }'
            'QMenu::item { padding: 4px 20px; }'
            'QMenu::item:selected { background-color: #2a5a8a; }'
        )
        sub = menu.addMenu('迁移到分组')
        for g in load_watchlist_groups():
            sub.addAction(g, lambda _=False, c=code, s=src_group, grp=g:
                          self._move_code_to_group(c, s, grp))
        menu.exec(self.table_watchlist.viewport().mapToGlobal(pos))

    def _move_code_to_group(self, code: str, src_group, target: str):
        rec = self._find_wl_record(code, src_group)
        if rec is None:
            return
        if (rec.get('group') or self._wl_default_group) == target:
            return
        tc = _code_to_tencent(code)
        clash = any(
            r is not rec
            and _code_to_tencent(r.get('code', '')) == tc
            and (r.get('group') or self._wl_default_group) == target
            for r in self._watchlist_codes
        )
        if clash:
            QMessageBox.information(self, '提示', f'分组「{target}」已存在该代码')
            return
        rec['group'] = target
        save_watchlist(self._watchlist_codes)
        self._rerender_watchlist()

    # ---- 自选股批量操作 ----

    def _wl_checked_targets(self) -> list[tuple]:
        """勾选行的 (code, 分组)；分组取自第0列 UserRole。"""
        out = []
        for r in range(self.table_watchlist.rowCount()):
            item = self.table_watchlist.item(r, 0)
            if item and item.checkState() == Qt.CheckState.Checked:
                c = _clean_code(item.text())
                if c:
                    out.append((c, item.data(Qt.ItemDataRole.UserRole)))
        return out

    def _wl_key(self, code, group):
        return (_code_to_tencent(code), group or self._wl_default_group)

    def _wl_checked_keys(self) -> set[tuple[str, str]]:
        return {
            self._wl_key(code, group)
            for code, group in self._wl_checked_targets()
        }

    def _on_wl_select_all(self, state):
        cs = Qt.CheckState.Checked if state == Qt.CheckState.Checked.value else Qt.CheckState.Unchecked
        for r in range(self.table_watchlist.rowCount()):
            item = self.table_watchlist.item(r, 0)
            if item:
                item.setCheckState(cs)

    def _reset_wl_select_all(self):
        self._wl_select_all.blockSignals(True)
        self._wl_select_all.setChecked(False)
        self._wl_select_all.blockSignals(False)

    def _on_wl_batch_migrate(self):
        targets = self._wl_checked_targets()
        if not targets:
            QMessageBox.information(self, '提示', '请先勾选股票')
            return
        groups = load_watchlist_groups()
        if not groups:
            QMessageBox.information(self, '提示', '暂无可迁移的分组')
            return
        target, ok = QInputDialog.getItem(
            self, '批量迁移', f'将选中的 {len(targets)} 只迁移到分组：', groups, 0, False)
        if not ok or not target:
            return
        default = self._wl_default_group
        # 目标组已有的代码（含本批已迁入的），用于跳过同组重复
        placed = {_code_to_tencent(r['code']) for r in self._watchlist_codes
                  if (r.get('group') or default) == target}
        moved = 0
        for code, group in targets:
            rec = self._find_wl_record(code, group)
            if rec is None or (rec.get('group') or default) == target:
                continue
            tc = _code_to_tencent(code)
            if tc in placed:
                continue
            rec['group'] = target
            placed.add(tc)
            moved += 1
        save_watchlist(self._watchlist_codes)
        self._reset_wl_select_all()
        self._rerender_watchlist()
        skipped = len(targets) - moved
        if skipped:
            QMessageBox.information(
                self, '批量迁移', f'迁移 {moved} 只，跳过 {skipped} 只（目标组已有该代码）。')

    def _on_wl_batch_delete(self):
        targets = self._wl_checked_targets()
        if not targets:
            QMessageBox.information(self, '提示', '请先勾选股票')
            return
        if QMessageBox.question(
                self, '批量删除', f'确认删除选中的 {len(targets)} 只自选股？',
                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No
        ) != QMessageBox.StandardButton.Yes:
            return
        default = self._wl_default_group
        keys = {(_code_to_tencent(c), g) for c, g in targets}
        self._watchlist_codes = [
            rec for rec in self._watchlist_codes
            if (_code_to_tencent(rec.get('code', '')), (rec.get('group') or default)) not in keys
        ]
        save_watchlist(self._watchlist_codes)
        self._reset_wl_select_all()
        self._rerender_watchlist()

    def _render_wl_row(self, r: int, row: dict, checked_keys: set[tuple[str, str]] | None = None):
            main_net = _to_float(row.get('main_net'))
            main_net_yi = main_net / 100000000 if abs(main_net) >= 1 else 0.0

            shares = row.get('shares')
            disp_shares = row.get('daily_shares')
            cum_pnl = row.get('cumulative_pnl')
            cum_pnl_pct = row.get('cumulative_pnl_pct')
            add_price = row.get('add_price')
            day_pnl = row.get('daily_pnl')
            shares_text = f'{disp_shares:,}' if disp_shares else '--'
            add_text = _format_price(add_price, row.get('code')) if add_price else '--'
            pre_close_text = _format_price(row.get('pre_close'), row.get('code')) if row.get('pre_close') else '--'
            pnl_text = f'{cum_pnl:+,.2f}' if cum_pnl is not None else '--'
            pnl_pct_text = f'{cum_pnl_pct:+.2f}%' if cum_pnl_pct is not None else '--'
            day_text = f'{day_pnl:+,.2f}' if day_pnl is not None else '--'

            values = [
                _clean_code(row['code']),
                row.get('name', ''),
                shares_text,
                add_text,
                _format_price(row.get('price'), row.get('code')),
                pre_close_text,
                _format_pct(row.get('pct'), signed=True),
                f'{main_net_yi:.2f}' if main_net_yi != 0 else '—',
                _format_pct(row.get('main_pct')) if row.get('main_pct') is not None else '—',
                _format_money(row.get('super_net')) if row.get('super_net') is not None else '—',
                _format_money(row.get('big_net')) if row.get('big_net') is not None else '—',
                pnl_text,
                pnl_pct_text,
                day_text,
            ]
            for c, text in enumerate(values):
                item = QTableWidgetItem(text)
                item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
                if c == 0:
                    item.setFlags(item.flags() | Qt.ItemFlag.ItemIsUserCheckable)
                    state = Qt.CheckState.Unchecked
                    if checked_keys and self._wl_key(row.get('code'), row.get('group')) in checked_keys:
                        state = Qt.CheckState.Checked
                    item.setCheckState(state)
                    item.setData(Qt.ItemDataRole.UserRole, row.get('group'))
                if c == 1:
                    item.setForeground(QBrush(QColor('#ffffff')))
                elif c == 6:  # 涨跌幅
                    item.setForeground(self._brush(row.get('pct')))
                elif 7 <= c <= 10:  # 资金流
                    val = [row.get('main_net'), row.get('main_pct'),
                           row.get('super_net'), row.get('big_net')][c - 7]
                    item.setForeground(self._brush(val))
                elif c in (11, 12):  # 累计盈亏 / 盈亏比例
                    if cum_pnl is not None and cum_pnl > 0:
                        item.setForeground(QBrush(QColor('#39c87a')))
                    elif cum_pnl is not None and cum_pnl < 0:
                        item.setForeground(QBrush(QColor('#e05c5c')))
                    else:
                        item.setForeground(QBrush(QColor('#aaaaaa')))
                elif c == 13:  # 当日收益
                    if day_pnl is not None and day_pnl > 0:
                        item.setForeground(QBrush(QColor('#39c87a')))
                    elif day_pnl is not None and day_pnl < 0:
                        item.setForeground(QBrush(QColor('#e05c5c')))
                    else:
                        item.setForeground(QBrush(QColor('#aaaaaa')))
                self.table_watchlist.setItem(r, c, item)

            # 操作按钮：编辑 + 删除
            ops = QWidget()
            ops.setStyleSheet(f'background-color: {CHART_BG};')
            ops_lay = QHBoxLayout(ops)
            ops_lay.setContentsMargins(2, 1, 2, 1)
            ops_lay.setSpacing(2)

            edit_btn = QPushButton('编辑')
            edit_btn.setFixedHeight(22)
            edit_btn.setStyleSheet(
                "QPushButton { background-color: #1a2a3a; color: #7acbe8; "
                "border: 1px solid #3a6a8a; border-radius: 3px; font-size: 10px; padding: 0 4px; }"
                "QPushButton:hover { background-color: #2a5a8a; color: #fff; }"
            )
            code = row['code']
            grp = row.get('group')
            edit_btn.clicked.connect(lambda checked, c=code, g=grp: self._edit_watchlist_record(c, g))
            ops_lay.addWidget(edit_btn)

            del_btn = QPushButton('删除')
            del_btn.setFixedHeight(22)
            del_btn.setStyleSheet(
                "QPushButton { background-color: #3a1a1a; color: #e84444; border: 1px solid #e84444; "
                "border-radius: 3px; font-size: 10px; padding: 0 4px; }"
                "QPushButton:hover { background-color: #e84444; color: #fff; }"
            )
            del_btn.clicked.connect(lambda checked, c=code, g=grp: self._remove_watchlist_code(c, g))
            ops_lay.addWidget(del_btn)

            self.table_watchlist.setCellWidget(r, len(_WATCHLIST_HEADERS) - 1, ops)
